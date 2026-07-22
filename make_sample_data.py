
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime

# ─────────────────────────────────────────
# 공통 설정
# ─────────────────────────────────────────
PARTNERS = ["삼성전자", "LG전자", "SK하이닉스", "현대모비스", "롯데케미칼"]
FONT_NAME = "Arial"

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
CELL_FONT   = Font(name=FONT_NAME, size=10)
EVEN_FILL   = PatternFill("solid", start_color="EBF3FB")
ODD_FILL    = PatternFill("solid", start_color="FFFFFF")

thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# v2.0 KT 컬럼명
KT_COLS = ["구매문서번호", "인수증", "입고승인일", "공급가액",
           "공급자명", "세금코드명", "이동유형명", "국책과제여부"]
KT_WIDTHS = [20, 14, 14, 16, 14, 12, 14, 14]

# v2.0 플랫폼 컬럼명
PL_COLS = ["주문번호", "입고일", "승인일", "일일정산일",
           "정산금액", "협력사명", "매출과세구분"]
PL_WIDTHS = [20, 14, 14, 14, 16, 14, 14]


def fmt_sheet(ws, cols, widths):
    for col_idx, (col_name, width) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 20


def write_row(ws, row_num, values, amt_col_idx):
    fill = EVEN_FILL if row_num % 2 == 0 else ODD_FILL
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=value)
        cell.font = CELL_FONT
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if col_idx == amt_col_idx:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right", vertical="center")


def d(day):
    return datetime.date(2025, 5, day)


# ─────────────────────────────────────────
# 데이터 정의
# ─────────────────────────────────────────

# [A] 일반 주문 15건 (KT ↔ 플랫폼 완전 일치)
# (구매문서번호, 인수증, 입고일day, 금액, 공급자명, 세금코드명)
normal_orders = [
    ("4500000001", "INV-001", 1,  1_200_000, "삼성전자",   "과세"),
    ("4500000002", "INV-002", 1,    850_000, "LG전자",     "과세"),
    ("4500000003", "INV-003", 2,  2_340_000, "SK하이닉스", "과세"),
    ("4500000004", "INV-004", 2,    670_000, "현대모비스", "면세"),
    ("4500000005", "INV-005", 3,  1_890_000, "롯데케미칼", "과세"),
    ("4500000006", "INV-006", 5,    530_000, "삼성전자",   "면세"),
    ("4500000007", "INV-007", 6,  3_100_000, "LG전자",     "과세"),
    ("4500000008", "INV-008", 7,    720_000, "SK하이닉스", "과세"),
    ("4500000009", "INV-009", 8,  1_450_000, "현대모비스", "과세"),
    ("4500000010", "INV-010", 9,    980_000, "롯데케미칼", "면세"),
    ("4500000011", "INV-011", 10, 2_650_000, "삼성전자",   "과세"),
    ("4500000012", "INV-012", 12,   440_000, "LG전자",     "면세"),
    ("4500000013", "INV-013", 13, 1_770_000, "SK하이닉스", "과세"),
    ("4500000014", "INV-014", 14,   610_000, "현대모비스", "과세"),
    ("4500000015", "INV-015", 15, 3_280_000, "롯데케미칼", "과세"),
]

# [B] KT에만 있는 주문 3건 (플랫폼 누락 — 핵심 테스트)
kt_only_orders = [
    ("4500000016", "INV-016", 4,    990_000, "삼성전자",   "과세"),
    ("4500000017", "INV-017", 8,  1_580_000, "LG전자",     "과세"),
    ("4500000018", "INV-018", 11,   730_000, "SK하이닉스", "면세"),
]

# [C] 반품 2건
# KT: 이동유형명 = "반품입고"
returns_kt = [
    ("4500000019", "INV-019", 6,  -450_000, "현대모비스", "과세"),
    ("4500000020", "INV-020", 13, -870_000, "롯데케미칼", "과세"),
]
# 플랫폼: 동일 주문번호로 반품
returns_platform = [
    ("4500000019", 4,  -450_000, "현대모비스", "과세"),
    ("4500000020", 11, -870_000, "롯데케미칼", "과세"),
]

# ─────────────────────────────────────────
# kt_raw.xlsx  (v2.0 컬럼)
# ─────────────────────────────────────────
wb_kt = openpyxl.Workbook()
ws_kt = wb_kt.active
ws_kt.title = "KT_입고내역"
ws_kt.freeze_panes = "A2"
fmt_sheet(ws_kt, KT_COLS, KT_WIDTHS)

all_kt_rows = []
for o in normal_orders:
    # (구매문서번호, 인수증, 입고승인일, 공급가액, 공급자명, 세금코드명, 이동유형명, 국책과제여부)
    all_kt_rows.append((o[0], o[1], d(o[2]), o[3], o[4], o[5], "일반입고", "N"))
for o in kt_only_orders:
    all_kt_rows.append((o[0], o[1], d(o[2]), o[3], o[4], o[5], "일반입고", "N"))
for o in returns_kt:
    all_kt_rows.append((o[0], o[1], d(o[2]), o[3], o[4], o[5], "반품입고", "N"))

for i, row_vals in enumerate(all_kt_rows, 2):
    write_row(ws_kt, i, row_vals, amt_col_idx=4)

wb_kt.save("kt_raw.xlsx")
print(f"[OK] kt_raw.xlsx 생성 완료 ({len(all_kt_rows)}건)")

# ─────────────────────────────────────────
# platform.xlsx  (v2.0 컬럼)
# ─────────────────────────────────────────
wb_pl = openpyxl.Workbook()
ws_pl = wb_pl.active
ws_pl.title = "플랫폼_입고내역"
ws_pl.freeze_panes = "A2"
fmt_sheet(ws_pl, PL_COLS, PL_WIDTHS)

all_pl_rows = []
for o in normal_orders:
    # (주문번호, 입고일, 승인일, 일일정산일, 정산금액, 협력사명, 매출과세구분)
    all_pl_rows.append((o[0], d(o[2]), d(o[2]), d(o[2]), o[3], o[4], o[5]))
for o in returns_platform:
    all_pl_rows.append((o[0], d(o[1]), d(o[1]), d(o[1]), o[2], o[3], o[4]))

for i, row_vals in enumerate(all_pl_rows, 2):
    write_row(ws_pl, i, row_vals, amt_col_idx=5)

wb_pl.save("platform.xlsx")
print(f"[OK] platform.xlsx 생성 완료 ({len(all_pl_rows)}건)")
print("\n[데이터 요약]")
print(f"  KT 전체    : {len(all_kt_rows)}건  (일반15 + KT전용3 + 반품2)")
print(f"  플랫폼 전체: {len(all_pl_rows)}건  (일반15 + 반품2)")
print(f"  KT전용(누락예상): {[o[0] for o in kt_only_orders]}")
print(f"  반품 구매문서번호: {[o[0] for o in returns_kt]}")
