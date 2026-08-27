# -*- coding: utf-8 -*-
"""배송 예측 정보 과제 - AX(외주사) 협의용 "AI 개선 여지 진단" 브리핑.

2026-07-27 신규: AX 프로젝트 외주사가 "배송예측 모델을 AI로 개선/보완할 여지가
있는지" 진단하겠다며 미팅을 요청 -> known_risks.md(내부 의사결정 로그, 외부인이
보기엔 맥락 없이 혼란스러움) 대신, 이 미팅 전용으로 별도 정리:
  ① 시스템 소개(통계적 접근이지 딥러닝이 아님을 명확히 - 불필요한 제안 예방)
  ② 이미 시도·검증·기각한 개선 아이디어(외주사가 같은 제안 반복 안 하도록)
  ③ 남은 한계와 그 중 AI가 실제로 도울 수 있는 지점(솔직하게 구분)
오프닝은 《더 퍼스트 미닛》(크리스 페닝)의 "의도(Intent) 먼저 밝히기" 적용.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1E, 0x27, 0x61)
DARKGRAY = RGBColor(0x36, 0x45, 0x4F)
MIDGRAY = RGBColor(0x6E, 0x7B, 0x8B)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xB0, 0x6A, 0x00)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_h1(doc, text):
    h = doc.add_paragraph()
    r = h.add_run(text)
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(5)
    return h


def add_body(doc, text, size=10, color=DARKGRAY, indent=0.2, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def make_table(doc, headers, rows, widths, center_from=None):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = w
    for c, w, lbl in zip(t.rows[0].cells, widths, headers):
        c.width = w
        set_cell_shading(c, "2A3673")
        pc = c.paragraphs[0]
        pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = pc.add_run(lbl)
        rc.font.bold = True
        rc.font.size = Pt(10)
        rc.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (row, values) in enumerate(zip(t.rows[1:], rows)):
        cells = row.cells
        shade = "F4F6FB" if i % 2 == 0 else "FFFFFF"
        for c, w in zip(cells, widths):
            c.width = w
            set_cell_shading(c, shade)
        for ci, (cell, val) in enumerate(zip(cells, values)):
            color, bold, center, size = DARKGRAY, False, (center_from is not None and ci >= center_from), 9.5
            text = val
            if isinstance(val, tuple):
                text, color = val
                bold, center = True, True
            p = cell.paragraphs[0]
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    return t


doc = Document("minimal_template.docx")

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, m, Cm(1.8))

style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(10.5)

# Title
p = doc.add_paragraph()
run = p.add_run("배송 예측 정보 과제")
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = NAVY
p.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
run = p.add_run("AI 개선 여지 진단 - AX 프로젝트 협의용  ·  2026-07-27  ·  네트워크사업1팀 이돈현")
run.font.size = Pt(10.5)
run.font.color.rgb = MIDGRAY
p.paragraph_format.space_after = Pt(8)

p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "1E2761")
pBdr.append(bottom)
pPr.append(pBdr)
p.paragraph_format.space_after = Pt(8)

# 오프닝(의도) - 더 퍼스트 미닛 프레이밍
box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
box.autofit = False
box.columns[0].width = Cm(17.4)
cell = box.rows[0].cells[0]
cell.width = Cm(17.4)
set_cell_shading(cell, "F4F6FB")
cell.paragraphs[0].paragraph_format.space_before = Pt(6)
cell.paragraphs[0].paragraph_format.space_after = Pt(2)
r = cell.paragraphs[0].add_run("오늘 미팅의 목적")
r.font.bold = True
r.font.size = Pt(11)
r.font.color.rgb = NAVY
p2 = cell.add_paragraph()
p2.paragraph_format.space_after = Pt(6)
r2 = p2.add_run(
    "이미 구축·검증을 마친 배송 예측 모델에, AI 기술로 추가 개선하거나 보완할 여지가 "
    "있는지 함께 진단하는 자리입니다. 지금까지 개선 아이디어 다수를 실측 데이터로 "
    "검증해왔기 때문에, 아래 순서로 공유드립니다: ① 무엇을 만들었는가 → ② 이미 시도하고 "
    "검증한 것 → ③ 아직 남은 한계와 AI가 실제로 도울 수 있는 지점."
)
r2.font.size = Pt(10.5)
r2.font.color.rgb = DARKGRAY

# ---------------------------------------------------------------
add_h1(doc, "① 지금 무엇을 만들었는가")

add_body(doc,
    "머신러닝/딥러닝 모델이 아니라, 과거 실제 배송 기록 265만 건을 통계적으로 집계한 "
    "\"확률표 + 자동 대체(폴백) 구조\"입니다. 상품마다 \"N일 이내 도착확률\"을 실측 비율로 "
    "그대로 계산하고, 자기 데이터가 부족한 상품은 상품 → 협력사×상품군 → 협력사 → "
    "표준납기일 → 전체 순으로 더 큰 표본의 평균을 자동으로 대신 씁니다.", after=4)
add_body(doc,
    "학습에 전혀 쓰지 않은 2026년 상반기 실제 주문 44.7만 건으로 검증한 결과, 예측확률과 "
    "실제 도착률의 차이(오차)는 전 구간 0.1~1.4%p 수준입니다.", after=8)

# ---------------------------------------------------------------
add_h1(doc, "② 이미 시도하고 검증한 개선 아이디어")
add_body(doc,
    "\"더 정교하게 나누거나 보정하면 좋아지지 않을까\"라는 아이디어들을 실제 홀드아웃 "
    "데이터로 검증했습니다. 결과는 대부분 \"효과 없음\" 또는 \"오히려 손해\"였습니다 - "
    "데이터가 이미 충분히 촘촘한 상태에서는 세분화가 표본을 쪼개 노이즈를 키우는 경우가 "
    "많았습니다.", after=6)

tried_rows = [
    ("레벨×day마크 편차 사전보정", ("기각", MIDGRAY), "편차 방향이 해마다 바뀌어 작년 기준으로 보정하면 올해는 더 틀림"),
    ("상품군 세분화(중분류→소분류→세분류)", ("기각", MIDGRAY), "일부(~20%)만 개선, 대다수(~80%)는 표본 쪼개져 악화 - 순손해"),
    ("최소표본 기준(MIN_N) 일괄 하향", ("조건부만 적용", AMBER), "일괄 하향은 과거 급성장 협력사 오판 사고 재발 위험 - 안전 확인된 일부만 완화"),
    ("공휴일 유형별 세부 보정 확대", ("부분 적용", AMBER), "표본이 다양해질수록(공휴일 종류가 섞일수록) 안전 구간이 오히려 좁아짐"),
    ("주문 요일(월~일) 반영", ("시도 후 롤백", MIDGRAY), "요일 자체 영향은 큼(최대 51%p 차이)이나 다른 요인과 섞으니 전체 정확도 악화"),
    ("표준납기일 구간 오차 개선(날짜분리·카테고리분리·학습기간 확장 등 4가지)", ("전부 기각", MIDGRAY), "네 각도 모두 same-order 검증에서 개선 확인 안 됨"),
]
make_table(doc, ["시도한 아이디어", "결과", "핵심 이유"], tried_rows, [Cm(6.0), Cm(2.6), Cm(8.8)])

add_body(doc,
    "※ \"same-order 검증\" = 이 프로젝트의 표준 검증법. 같은 주문 집합을 놓고 \"기존 방식으로 "
    "예측했을 때\"와 \"새 방식으로 예측했을 때\"를 나란히 비교해, 새 방식이 실제로 더 정확할 "
    "때만 채택합니다. ※ \"요일 반영\" 사례처럼 어떤 요인 하나의 영향력이 커도(51%p), 이미 "
    "쓰고 있는 다른 분류(협력사×상품군 등)와 같이 섞으면 표본이 더 잘게 쪼개져 각 칸의 데이터가 "
    "부족해지고, 그 노이즈가 원래 효과보다 더 커져서 전체 정확도가 떨어지는 경우가 많았습니다.",
    size=9, color=MIDGRAY, after=8)

# ---------------------------------------------------------------
add_h1(doc, "③ 아직 남은 한계 - AI가 실제로 도울 수 있는 지점")
add_body(doc,
    "모든 한계가 AI로 풀리는 건 아닙니다. 아래는 원인 성격별로 구분한 것입니다 - "
    "\"예측 불가능한 사건\"과 \"더 나은 기법으로 풀릴 수 있는 문제\"는 다릅니다.", after=6)

limit_rows = [
    ("특정 협력사의 일회성 지연\n(예: 한 협력사가 특정 한 달만 급격히 느려짐)",
     "예측 시점엔 어느 주문이 그 이상치에 해당할지 알 방법이 없어 사전 예측은 원천적으로 불가",
     "가능성 있음 - 현재 월간 규칙기반(임계값) 감시를 이상탐지(anomaly detection) 모델로 고도화하면 오탐/누락 감소 여지"),
    ("낙찰 협력사 전환 자동 감지",
     "협력사가 바뀌어도 sku 레벨 정확도 자체는 저하 안 됨(실측 확인) - 다만 \"전환 시점\"을 "
     "자동으로 알아채는 기능은 아직 없음",
     "가능성 있음 - 미구현 상태. 협력사 이력 변화를 자동 감지하는 모니터링 추가 가능"),
    ("표준납기일 기반 예측(협력사 이력 얕은 상품)의 4~7일 구간 잔존 오차",
     "데이터 부족이 아니라 이미 확인된 한계 - 학습기간을 2022년, 2018~2021년(근사값)까지 "
     "확장하는 시도를 포함 4가지 각도 모두 검증했고, 그 결과가 세 차례 일관되게 \"개선 없음\"으로 "
     "나옴(IT 데이터 추가요청도 검토 후 철회)",
     "낮음 - \"데이터를 더 모으면 풀린다\"는 이미 배제됨. 재고·생산 상태처럼 지금 없는 "
     "새로운 종류의 정보가 있어야 풀릴 가능성이 높음"),
    ("표준납기일이 유독 긴 상품의 원인 미파악\n(예: 영풍기획 물티슈, 15~28일)",
     "상품명·기존 데이터만으로는 원인 확인 불가 - 라벨·스티커 등 제작 공정이 필요한 상품인지가 "
     "어디에도 구조화돼 있지 않음. 상품명에 관련 키워드(스티커·라벨·각인 등)가 있는 상품 2,024개를 "
     "확인해봐도 오히려 표준납기가 더 짧아(평균 8.8일 vs 14.2일), 상품명만으로는 판단 불가",
     "가능성 있음 - 상품명·협력사 정보를 보고 \"제작 공정이 필요한 상품인가\"를 판단해 정보를 "
     "채워주는 에이전트를 붙이면, 표준납기일이 긴 이유를 설명하고 예측에도 반영할 수 있는 "
     "새로운 정보 축이 될 수 있음"),
]
make_table(doc, ["남은 한계", "원인", "AI가 도울 여지"], limit_rows, [Cm(4.6), Cm(6.4), Cm(6.4)])

add_body(doc,
    "정리하면, \"모델을 더 정교하게 만드는 방향\"은 이미 상당 부분 검증을 마쳤고 추가 여지가 "
    "크지 않습니다. AI가 실제로 기여할 수 있는 지점은 오히려 \"이상 감지·모니터링 자동화\"와 "
    "\"현재 데이터에 없는 정보를 찾아 채워주는 에이전트\" 두 방향에 가깝습니다.",
    size=10, color=DARKGRAY, after=4)

doc.save("배송예측_AX외주사_AI개선여지진단.docx")
print("saved")
