# -*- coding: utf-8 -*-
"""배송 예측 정보 과제 - 1p 모델 소개 보고서.

2026-07-14 재구성: "검토 미팅 후속 조치 보고" -> 미팅에 없었던 임직원에게 재배포해도
바로 이해되는 "모델을 처음 소개하는 문서"로 전환(사용자 요청). 미팅 맥락이 필요한
내용(지시사항 검증표 등)은 build_change_log.py 산출물로 위임.

2026-07-15 스타일 개편: 사용자가 본부장 보고에 실제로 썼던 편집본("1. 배송 예측 정보
과제 현황 보고_V2.pdf" = "배송예측 보고_편집본.docx")의 형식을 따름 - ▣ 섹션 마커,
• 개조식 불릿(서술문 대신 압축 구문), -> 로 결과 연결, 핵심 결론 박스 + 성과 지표 표 +
다음 단계(담당 협조) 구조. 원본 docx는 사내 DRM(SCDSA004)이라 프로그램으로 못 읽고,
PDF에서 텍스트를 수동 추출(zlib+ToUnicode 파싱)해서 구조를 파악함."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1E, 0x27, 0x61)
DARKGRAY = RGBColor(0x36, 0x45, 0x4F)
MIDGRAY = RGBColor(0x6E, 0x7B, 0x8B)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xB0, 0x6A, 0x00)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def tighten_table(table, top=20, bottom=20, left=60, right=60):
    """표 전체의 셀 여백을 촘촘하게(twips 단위) - 문단 간격만으로는 못 줄이는
    셀 기본 여백까지 줄여서 표 여러 개가 한 페이지에 들어가게 함."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tblPr.append(mar)


def add_section(doc, text):
    h = doc.add_paragraph()
    r = h.add_run("▣ " + text)
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY
    h.paragraph_format.space_before = Pt(4)
    h.paragraph_format.space_after = Pt(1)
    return h


def add_bullet(doc, text, size=10, indent=0.35, after=3, color=DARKGRAY):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run("• " + text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def make_table(doc, headers, rows, widths):
    """텍스트 불릿 대신 스캔하기 쉬운 표. headers 없이(빈 리스트) 쓰면 헤더행 생략."""
    n_header = 1 if headers else 0
    t = doc.add_table(rows=len(rows) + n_header, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tighten_table(t)
    for col, w in zip(t.columns, widths):
        col.width = w
    if headers:
        for c, w, lbl in zip(t.rows[0].cells, widths, headers):
            c.width = w
            set_cell_shading(c, "2A3673")
            pc = c.paragraphs[0]
            pc.paragraph_format.space_before = Pt(2)
            pc.paragraph_format.space_after = Pt(2)
            rc = pc.add_run(lbl)
            rc.font.bold = True
            rc.font.size = Pt(9.5)
            rc.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (row, values) in enumerate(zip(t.rows[n_header:], rows)):
        cells = row.cells
        shade = "F4F6FB" if i % 2 == 0 else "FFFFFF"
        for c, w in zip(cells, widths):
            c.width = w
            set_cell_shading(c, shade)
        for ci, (cell, val) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(9)
            r.font.bold = (ci == 0)
            r.font.color.rgb = NAVY if ci == 0 else DARKGRAY
    return t


def make_flow(doc, steps, widths):
    """문제->원인->해결 같은 3~4단계 흐름을 화살표 없이 색상 대비로 표현하는 가로 도식."""
    t = doc.add_table(rows=2, cols=len(steps))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tighten_table(t)
    for col, w in zip(t.columns, widths):
        col.width = w
    for c, w, (tag, _, color) in zip(t.rows[0].cells, widths, steps):
        c.width = w
        set_cell_shading(c, color)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(tag)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for c, w, (_, desc, _) in zip(t.rows[1].cells, widths, steps):
        c.width = w
        set_cell_shading(c, "F4F6FB")
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(desc)
        r.font.size = Pt(8.5)
        r.font.color.rgb = DARKGRAY
    return t


doc = Document("minimal_template.docx")

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, m, Cm(0.9))

style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(10.5)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = p.add_run("배송 예측 정보 과제")
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = NAVY
p.paragraph_format.space_after = Pt(1)

p = doc.add_paragraph()
run = p.add_run("모델 소개 및 신뢰성 검증  ·  2026-07-15  ·  네트워크사업1팀 이돈현")
run.font.size = Pt(11)
run.font.color.rgb = MIDGRAY
p.paragraph_format.space_after = Pt(6)

# divider
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "1E2761")
pBdr.append(bottom)
pPr.append(pBdr)
p.paragraph_format.space_after = Pt(4)

# Key conclusion callout
box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
box.autofit = False
tighten_table(box)
box.columns[0].width = Cm(19.0)
cell = box.rows[0].cells[0]
cell.width = Cm(19.0)
set_cell_shading(cell, "F4F6FB")

# 오프닝 문장(2026-07-27 추가) - 《더 퍼스트 미닛》(크리스 페닝)의 "의도(Intent)를
# 맨 먼저 밝힌다" 원칙 적용. 별도 문단 대신 박스 첫 줄로 넣어 공간을 아낌(1p 여백 한계).
cell.paragraphs[0].paragraph_format.space_before = Pt(2)
cell.paragraphs[0].paragraph_format.space_after = Pt(0)
r_intent = cell.paragraphs[0].add_run("※ 승인 요청이 아닌 진행상황 공유 - ⑤단계(통합플랫폼 적용) 직전입니다.")
r_intent.font.size = Pt(9)
r_intent.font.italic = True
r_intent.font.color.rgb = MIDGRAY

p_title = cell.add_paragraph()
p_title.paragraph_format.space_before = Pt(2)
p_title.paragraph_format.space_after = Pt(1)
r = p_title.add_run("핵심 결론")
r.font.bold = True
r.font.size = Pt(11)
r.font.color.rgb = NAVY
p2 = cell.add_paragraph()
p2.paragraph_format.space_after = Pt(4)
r2 = p2.add_run("과거 실제 배송 기록을 통계적으로 계산해, 고객이 주문 전에 “N일 안에 도착할 확률 "
                "X%”를 미리 알 수 있게 하는 예측 모델을 개발 중입니다. 학습에 전혀 쓰지 않은 "
                "2026년 상반기 실제 주문 44.7만 건으로 검증한 결과, 예측과 실제의 차이(오차)는 "
                "0.1~1.4%p 수준입니다.")
r2.font.size = Pt(11)
r2.font.color.rgb = DARKGRAY

# ---------------------------------------------------------------
# 진행 단계(2026-07-24 추가): 본부장이 정의한 5단계 로드맵 그대로 시각화 -
# "지금 어디까지 왔고 다음이 뭔가"를 첫 화면에서 바로 답하기 위함.
stage_label = doc.add_paragraph()
stage_label.paragraph_format.space_before = Pt(4)
stage_label.paragraph_format.space_after = Pt(2)
r_sl = stage_label.add_run("현재 진행 단계")
r_sl.font.bold = True
r_sl.font.size = Pt(9.5)
r_sl.font.color.rgb = MIDGRAY

stages = [
    ("①", "개념 수립", True),
    ("②", "기준 정의", True),
    ("③", "운영방안 설계", True),
    ("④", "신뢰도 검증", True),
    ("⑤", "데이터 적용(표시)", False),
]
stage_table = doc.add_table(rows=1, cols=len(stages))
stage_table.alignment = WD_TABLE_ALIGNMENT.CENTER
stage_table.autofit = False
tighten_table(stage_table)
sw2 = Cm(3.8)
for col in stage_table.columns:
    col.width = sw2
for cell, (num, label, done) in zip(stage_table.rows[0].cells, stages):
    cell.width = sw2
    set_cell_shading(cell, "2E7D32" if done else "B06A00")
    p0 = cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(2)
    r0 = p0.add_run(num + " " + label + ("  완료" if done else "  진행중(직전)"))
    r0.font.bold = True
    r0.font.size = Pt(9)
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p0.paragraph_format.space_after = Pt(2)

# ---------------------------------------------------------------
# 두괄식(2026-07-15): 결론 바로 아래에 핵심 근거(숫자)를 배치 - 위쪽 1/3만 읽어도
# "무엇을 만들었고 얼마나 정확한가"가 완결되도록. 설명(하는 일)과 방법론(검증)은 그 아래.
add_section(doc, "핵심 성과 지표")

stats = [
    ("0.1~1.4%p", "예측 오차", "“N일까지 도착확률” 예측치와 2026년 상반기 실제 도착률의 차이 (학습에 쓰지 않은 44.7만 건으로 검증)"),
    ("3,077개", "학습된 세그먼트", "상품코드 → 협력사×상품군 → 협력사 → 표준납기 → 상품종류 → 전체 순으로 촘촘하게 대체 (상품코드 2,171개 포함)"),
    ("265만 건", "학습 데이터 규모", "2023~2025년 4개 유통채널(KT/그룹사/외부사/지입자재) 전체 주문 (협력사×상품군 단위는 2018년까지 확장해 601만 건 사용)"),
]

table = doc.add_table(rows=len(stats), cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
tighten_table(table)
widths = [Cm(3.6), Cm(4.5), Cm(10.9)]
for col, w in zip(table.columns, widths):
    col.width = w
for row, (num, label, desc) in zip(table.rows, stats):
    cells = row.cells
    for c, w in zip(cells, widths):
        c.width = w
    set_cell_shading(cells[0], "1E2761")
    p_num = cells[0].paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_num.paragraph_format.space_before = Pt(2)
    p_num.paragraph_format.space_after = Pt(2)
    r_num = p_num.add_run(num)
    r_num.font.bold = True
    r_num.font.size = Pt(13)
    r_num.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    set_cell_shading(cells[1], "F4F6FB")
    cells[1].paragraphs[0].paragraph_format.space_before = Pt(2)
    cells[1].paragraphs[0].paragraph_format.space_after = Pt(2)
    r_label = cells[1].paragraphs[0].add_run(label)
    r_label.font.bold = True
    r_label.font.size = Pt(10.5)
    r_label.font.color.rgb = NAVY

    set_cell_shading(cells[2], "F4F6FB")
    cells[2].paragraphs[0].paragraph_format.space_before = Pt(2)
    cells[2].paragraphs[0].paragraph_format.space_after = Pt(2)
    r_desc = cells[2].paragraphs[0].add_run(desc)
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = DARKGRAY

# ---------------------------------------------------------------
add_section(doc, "이 모델이 하는 일")

how_rows = [
    ("예측 형태", "“오늘 주문하면 N일 이내 도착확률 X%” (네이버쇼핑 “7/5까지 도착 확률 92%”와 동일 방식 - "
                 "그날까지 누적 도착확률, 더 일찍 도착도 포함)"),
    ("계산 방식", "그 상품이 과거 실제 며칠 만에 도착했는지를 그대로 “N일 이내 도착 비율”로 사용 "
                 "(실제 배송 데이터 265만 건 근거, 계산 과정 전체 검증 가능)"),
    ("기록이 부족한 상품", "상품 → 협력사 → 비슷한 상품 종류 → 전체 평균 순으로 표본이 더 큰 집단의 "
                        "평균으로 자동 대체 (적은 표본의 우연한 숫자보다 안전)"),
    ("신규 협력사(거래 기록 없음)", "계약상 “표준납기일”은 항상 존재 → 비슷한 납기의 다른 상품 실제 기록으로 대신 예측"),
]
make_table(doc, ["구분", "방식"], how_rows, [Cm(4.4), Cm(14.6)])

# ---------------------------------------------------------------
add_section(doc, "정확도 검증 방법 - “모의고사” 방식 (학습에 안 쓴 미래 데이터로 채점)")

exam_rows = [
    ("학습 데이터", "2023~2025년 주문 265만 건"),
    ("검증(채점) 데이터", "2026년 상반기 실제 주문 44.7만 건 (학습에 전혀 미사용)"),
    ("채점 기준", "예측확률 vs 실제 도착비율 대조 (예: 예측 80%·실제 82% → 오차 2%p), "
                "9개 지점(1·2·3·5·7·10·14·21·30일)에서 반복 채점"),
    ("이 시험이 잡아낸 문제(개발 초기)", "최대 12.7%p 오차 → 원인: 급증 협력사를 “옛날 소량 기록”만 믿고 계산 → "
                                    "해결: 자기 기록 사용 기준 30건→1,000건 상향, 이후 같은 시험 통과분만 반영"),
]
make_table(doc, ["구분", "내용"], exam_rows, [Cm(3.6), Cm(15.4)])

# ---------------------------------------------------------------
step_p = doc.add_paragraph()
step_p.paragraph_format.space_before = Pt(3)
step_p.paragraph_format.space_after = Pt(0)
r_step_tag = step_p.add_run("▣ 다음 단계: ")
r_step_tag.font.bold = True
r_step_tag.font.size = Pt(10.5)
r_step_tag.font.color.rgb = NAVY
r_step_desc = step_p.add_run("통합플랫폼 서빙 연동 협의(박형근 차장)")
r_step_desc.font.size = Pt(9.5)
r_step_desc.font.color.rgb = DARKGRAY
r_foot = step_p.add_run("  ※ 상세는 별첨 참고")
r_foot.font.size = Pt(8)
r_foot.font.italic = True
r_foot.font.color.rgb = MIDGRAY

# 최종 패스: 모든 문단(표 안 포함) 줄간격을 명시적으로 단일行(1.0)로 고정 -
# 템플릿 상속 배수 줄간격이 숨어 있으면 여백 조정이 안 먹히는 원인이 되므로 확정.
def _force_single_spacing(container):
    for p in container.paragraphs:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for c in row.cells:
                _force_single_spacing(c)

_force_single_spacing(doc)

doc.save("배송예측_본부장보고_1p.docx")
print("saved")
