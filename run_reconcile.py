
"""
정산 대사 스크립트
────────────────────────────────────────────────
입력: kt_raw.xlsx / platform.xlsx
출력: 정산_대사결과.xlsx
  - Sheet1 (일치내역)  : KT ↔ 플랫폼 주문번호 기준 매칭된 건
  - Sheet2 (KT누락)    : KT에 있으나 플랫폼에 없는 건  ← 3건이어야 정상
  - Sheet3 (반품대사)  : 반품 주문을 요청번호로 매칭한 건 ← 2건이어야 정상
"""

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

COLORS = {
    "matched":  "1F4E79",   # 파랑  - 일치 헤더
    "missing":  "C00000",   # 빨강  - KT 누락 헤더
    "returns":  "375623",   # 초록  - 반품 헤더
    "row_even_matched":  "EBF3FB",
    "row_even_missing":  "FCE4D6",
    "row_even_returns":  "EBF1DE",
}

thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

BASE_COLS = ["주문번호", "요청번호", "입고일", "정산금액", "협력사명", "매출과세구분", "주문유형"]
COL_WIDTHS = [22, 18, 13, 16, 14, 14, 10]

def style_header(ws, col_count, header_color):
    hdr_fill = PatternFill("solid", start_color=header_color)
    hdr_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
    for c in range(1, col_count + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 20

def write_df_to_sheet(ws, df, header_color, even_color):
    cols = list(df.columns)
    # 헤더
    for c, col_name in enumerate(cols, 1):
        ws.cell(row=1, column=c, value=col_name)
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS[c-1] if c <= len(COL_WIDTHS) else 15
    style_header(ws, len(cols), header_color)

    # 데이터 행
    for r_idx, (_, row) in enumerate(df.iterrows(), 2):
        fill_color = even_color if r_idx % 2 == 0 else "FFFFFF"
        fill = PatternFill("solid", start_color=fill_color)
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cols[c_idx-1] == "정산금액":
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.freeze_panes = "A2"

# ─────────────────────────────────────────
# 1. 파일 읽기
# ─────────────────────────────────────────
print("📂 파일 읽는 중...")
df_kt = pd.read_excel("kt_raw.xlsx", dtype={"주문번호": str, "요청번호": str})
df_pl = pd.read_excel("platform.xlsx", dtype={"주문번호": str, "요청번호": str})

print(f"   KT     : {len(df_kt)}건")
print(f"   플랫폼 : {len(df_pl)}건")

# ─────────────────────────────────────────
# 2. 대사 로직
# ─────────────────────────────────────────

# (A) 반품 분리
kt_normal  = df_kt[df_kt["주문유형"] == "일반"].copy()
kt_return  = df_kt[df_kt["주문유형"] == "반품"].copy()

pl_normal  = df_pl[df_pl["주문유형"] == "일반"].copy()
pl_return  = df_pl[df_pl["주문유형"] == "반품"].copy()

# (B) 일반 주문 주문번호 기준 매칭
pl_order_nos = set(pl_normal["주문번호"])
matched_mask   = kt_normal["주문번호"].isin(pl_order_nos)
df_matched     = kt_normal[matched_mask].copy()       # Sheet1
df_kt_missing  = kt_normal[~matched_mask].copy()      # Sheet2 (KT에만 존재)

# (C) 반품 요청번호 기준 매칭
#   KT 반품(RTN-...) 의 요청번호  ↔  플랫폼 반품(ORD-...) 의 요청번호
df_return_kt = kt_return.rename(columns={c: f"KT_{c}" for c in kt_return.columns})
df_return_pl = pl_return.rename(columns={c: f"PL_{c}" for c in pl_return.columns})

df_return_matched = pd.merge(
    df_return_kt,
    df_return_pl,
    left_on="KT_요청번호",
    right_on="PL_요청번호",
    how="inner"
)

# 출력용 컬럼 정리
return_cols = [
    "KT_주문번호", "KT_요청번호", "KT_입고일", "KT_정산금액", "KT_협력사명",
    "PL_주문번호", "PL_입고일",   "PL_정산금액",
]
df_return_out = df_return_matched[return_cols].copy()
df_return_out.columns = [
    "KT_주문번호(반품)", "요청번호", "KT_입고일", "KT_정산금액", "협력사명",
    "PL_원주문번호",    "PL_입고일", "PL_정산금액",
]

# ─────────────────────────────────────────
# 3. 결과 출력
# ─────────────────────────────────────────
print("\n📊 대사 결과")
print(f"  ✅ Sheet1 (일치내역)  : {len(df_matched)}건")
print(f"  ⚠️  Sheet2 (KT누락)   : {len(df_kt_missing)}건  {'← 정상 (3건)' if len(df_kt_missing)==3 else '← ❌ 예상과 다름!'}")
print(f"  🔄 Sheet3 (반품대사)  : {len(df_return_out)}건  {'← 정상 (2건)' if len(df_return_out)==2 else '← ❌ 예상과 다름!'}")

if len(df_kt_missing) > 0:
    print("\n  [KT누락 주문번호]")
    for v in df_kt_missing["주문번호"].tolist():
        print(f"    - {v}")

if len(df_return_out) > 0:
    print("\n  [반품 매칭 내역]")
    for _, row in df_return_out.iterrows():
        print(f"    KT:{row['KT_주문번호(반품)']}  ↔  PL:{row['PL_원주문번호']}  (요청번호:{row['요청번호']})")

# ─────────────────────────────────────────
# 4. 결과 Excel 저장
# ─────────────────────────────────────────
out_path = "정산_대사결과.xlsx"
wb = openpyxl.Workbook()

# Sheet1: 일치내역
ws1 = wb.active
ws1.title = "①일치내역"
write_df_to_sheet(ws1, df_matched[BASE_COLS], COLORS["matched"], COLORS["row_even_matched"])

# Sheet2: KT누락
ws2 = wb.create_sheet("②KT누락")
write_df_to_sheet(ws2, df_kt_missing[BASE_COLS], COLORS["missing"], COLORS["row_even_missing"])

# Sheet3: 반품대사
ws3 = wb.create_sheet("③반품대사")
return_display_cols = list(df_return_out.columns)
# 반품 시트 열 너비 재정의
ret_widths = [22, 18, 13, 16, 14, 22, 13, 16]
for c_idx, width in enumerate(ret_widths, 1):
    ws3.column_dimensions[get_column_letter(c_idx)].width = width

write_df_to_sheet(ws3, df_return_out, COLORS["returns"], COLORS["row_even_returns"])

wb.save(out_path)
print(f"\n💾 결과 저장 완료: {out_path}")

# ─────────────────────────────────────────
# 5. 최종 검증
# ─────────────────────────────────────────
ok = True
if len(df_kt_missing) != 3:
    print("❌ [검증 실패] Sheet2(KT누락) 건수가 3건이어야 합니다.")
    ok = False
if len(df_return_out) != 2:
    print("❌ [검증 실패] Sheet3(반품대사) 건수가 2건이어야 합니다.")
    ok = False
if ok:
    print("\n🎉 모든 검증 통과! 대사 로직이 정상 작동합니다.")
