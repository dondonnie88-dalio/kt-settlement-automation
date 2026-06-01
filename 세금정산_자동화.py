# -*- coding: utf-8 -*-
"""
세금 정산 자동화 스크립트
입력 : 매출.xlsx  (시트: 2024년, 2025년)
출력 : 정산보고서_YYYY-MM.xlsx  (실행 월 자동 반영)

[1단계] 세금 구분별 부가세 정합성 검증
[2단계] 어음 발행 대상 판별 (비과세 전용)
[3단계] 4개 시트 Excel 보고서 생성
"""

import sys, os, math
from pathlib import Path
from datetime import datetime

for pkg in ["pandas", "openpyxl"]:
    try: __import__(pkg)
    except ImportError:
        print(f"[설치 필요] pip install {pkg}"); sys.exit(1)

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings
warnings.filterwarnings("ignore")

# ════════════════════════════════════════════════════════════════════════════
# 설정값 (필요 시 수정)
# ════════════════════════════════════════════════════════════════════════════
BASE_DIR    = Path(__file__).parent
INPUT_FILE  = BASE_DIR / "매출.xlsx"
기준월      = datetime.today().strftime("%Y-%m")
OUTPUT_FILE = BASE_DIR / f"정산보고서_{기준월}.xlsx"

VAT_오차허용   = 1          # ±1원 이내는 정상 처리
어음_최소금액  = 200_000_000
어음_대상규모  = {"대기업", "중견기업"}

# 데이터에 '거래처규모' 컬럼이 없을 경우 여기서 직접 설정
거래처규모_매핑 = {
    "오렌지스펙트럼(주)": "중소기업",
    "(주)케이티링커스":   "중견기업",
    "엔씨소프트(주)":    "대기업",
    "(주)다우기술":      "중소기업",
    "삼성SDS(주)":      "대기업",
    "LG CNS(주)":       "대기업",
}

# ════════════════════════════════════════════════════════════════════════════
# 스타일 헬퍼
# ════════════════════════════════════════════════════════════════════════════
FONT = "맑은 고딕"

BG = dict(
    header  = "1F4E79", title   = "2E75B6",
    error   = "FF0000", error_l = "FFCCCC",
    ok      = "375623", ok_l    = "C6EFCE",
    warn    = "7F6000", warn_l  = "FFE699",
    bill    = "1F3864", bill_l  = "BDD7EE",
    total   = "FFF2CC", alt     = "DEEAF1",
    grey    = "F2F2F2",
)
FG = dict(white="FFFFFF", red="C00000", green="375623",
          warn="7F6000", dark="1F4E79")
NUM   = "#,##0"
NUM_S = '#,##0;[Red]-#,##0;"-"'
PCT   = "0.00%"

def _s(): return Side(style="thin", color="BFBFBF")
def BD(): return Border(left=_s(), right=_s(), top=_s(), bottom=_s())

def hc(cell, text, bg=BG["header"], fg=FG["white"], bold=True, wrap=False, size=10):
    cell.value = text
    cell.font  = Font(name=FONT, bold=bold, color=fg, size=size)
    cell.fill  = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=wrap)
    cell.border = BD()

def dc(cell, value, fmt=None, bold=False, bg=None, align="right",
       color="000000"):
    cell.value = value
    cell.font  = Font(name=FONT, bold=bold, color=color, size=10)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BD()
    if fmt: cell.number_format = fmt
    if bg:  cell.fill = PatternFill("solid", start_color=bg)

def cw(ws, col, w): ws.column_dimensions[get_column_letter(col)].width = w

def sheet_title(ws, t1, t2=""):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=22)
    c = ws.cell(1, 1, t1)
    c.font = Font(name=FONT, bold=True, size=14, color=FG["white"])
    c.fill = PatternFill("solid", start_color=BG["title"])
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    if t2:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=22)
        c2 = ws.cell(2, 1, t2)
        c2.font = Font(name=FONT, size=10, color="595959")
        c2.fill = PatternFill("solid", start_color="D6E4F0")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

def total_row(ws, row, n_cols, skip_cols=None):
    """합계 행 공식 삽입"""
    skip_cols = skip_cols or set()
    for c in range(1, n_cols + 1):
        cell = ws.cell(row, c)
        cell.fill = PatternFill("solid", start_color=BG["total"])
        cell.border = BD()
        cell.font   = Font(name=FONT, bold=True)
        cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row, 1, "합  계").alignment = Alignment(
        horizontal="center", vertical="center")


# ════════════════════════════════════════════════════════════════════════════
# 1. 데이터 로드
# ════════════════════════════════════════════════════════════════════════════
def load_data():
    if not INPUT_FILE.exists():
        sys.exit(f"[오류] 파일 없음: {INPUT_FILE}")

    xl = pd.ExcelFile(INPUT_FILE)
    print(f"[읽기] 시트: {xl.sheet_names}")

    frames = []
    for sheet in xl.sheet_names:
        df = xl.parse(sheet)
        df["_시트"] = sheet
        frames.append(df)
        print(f"  [{sheet}] {len(df):,}행")

    data = pd.concat(frames, ignore_index=True)

    # 필수 컬럼
    must = ["협력사명", "주문번호", "매출과세구분",
            "정산수량", "정산단가", "정산금액"]
    missing = [c for c in must if c not in data.columns]
    if missing:
        sys.exit(f"[오류] 필수 컬럼 없음: {missing}")

    # 숫자 변환
    for col in ["정산수량", "정산단가", "주문금액", "정산금액",
                "매입금액", "매출총이익"]:
        if col in data.columns:
            data[col] = pd.to_numeric(data[col], errors="coerce").fillna(0)

    # 거래처규모 컬럼 처리
    if "거래처규모" not in data.columns:
        data["거래처규모"] = data["협력사명"].map(거래처규모_매핑).fillna("미분류")
    else:
        data["거래처규모"] = data["거래처규모"].fillna(
            data["협력사명"].map(거래처규모_매핑).fillna("미분류"))

    print(f"\n[전체] {len(data):,}행 | 과세구분: "
          f"{data['매출과세구분'].value_counts().to_dict()}")
    return data


# ════════════════════════════════════════════════════════════════════════════
# 2. [1단계] 부가세 정합성 검증
# ════════════════════════════════════════════════════════════════════════════
def validate_vat(data: pd.DataFrame) -> pd.DataFrame:
    df = data.copy()

    # 공급가액 = 정산수량 × 정산단가  (세금 제외 기준가)
    df["공급가액"] = df["정산수량"] * df["정산단가"]
    df["세금구분"]  = df["매출과세구분"].fillna("미확인").str.strip()

    결과 = []
    for _, row in df.iterrows():
        구분  = row["세금구분"]
        공가  = row["공급가액"]
        실제  = row["정산금액"]   # 실제 청구액

        if 구분 == "과세":
            부가세_계산  = round(공가 * 0.1)
            올바른청구액 = 공가 + 부가세_계산
        else:
            부가세_계산  = 0
            올바른청구액 = 공가   # 면세·비과세 = 공급가액 그대로

        오차금액 = 실제 - 올바른청구액
        결과.append({
            "부가세":      부가세_계산,
            "올바른청구액": 올바른청구액,
            "오차금액":    오차금액,
            "오차여부":    abs(오차금액) > VAT_오차허용,
        })

    계산컬럼 = pd.DataFrame(결과, index=df.index)
    df = pd.concat([df, 계산컬럼], axis=1)

    오차수 = df["오차여부"].sum()
    print(f"[1단계] 부가세 검증 완료 | 총 {len(df):,}행 | "
          f"오차 {오차수}건 / 정상 {len(df)-오차수}건")
    return df


# ════════════════════════════════════════════════════════════════════════════
# 3. [2단계] 어음 발행 대상 판별
# ════════════════════════════════════════════════════════════════════════════
def determine_bill(df: pd.DataFrame) -> pd.DataFrame:
    # Step A: 비과세만
    비과세 = df[df["세금구분"] == "비과세"].copy()
    if len(비과세) == 0:
        print("[2단계] 비과세 항목 없음 → 어음 발행 대상 없음")
        return pd.DataFrame()

    # Step B: (주문번호, 협력사명) 기준 합산
    grp = 비과세.groupby(["주문번호", "협력사명"], sort=False).agg(
        합산청구액  = ("정산금액",  "sum"),
        오차건수    = ("오차여부",  "sum"),
        거래처규모  = ("거래처규모","first"),
        사업장      = ("사업장",    "first"),
        정산확정일  = ("정산확정일","first"),
        품목수      = ("정산금액",  "count"),
    ).reset_index()

    # Step C: 3가지 조건 판별
    def classify(row):
        금액OK  = row["합산청구액"] >= 어음_최소금액
        규모OK  = row["거래처규모"] in 어음_대상규모
        오차없음 = row["오차건수"] == 0

        if not 금액OK or not 규모OK:
            return None          # 목록에서 제외
        if not 오차없음:
            return "오차 확인 필요"
        return "어음 발행 대상"

    grp["판별결과"] = grp.apply(classify, axis=1)
    result = grp[grp["판별결과"].notna()].copy()
    result.sort_values("합산청구액", ascending=False, inplace=True)

    발행대상 = (result["판별결과"] == "어음 발행 대상").sum()
    오차대상  = (result["판별결과"] == "오차 확인 필요").sum()
    print(f"[2단계] 어음 발행 판별 완료 | "
          f"발행 대상 {발행대상}건 | 오차 확인 필요 {오차대상}건")
    return result


# ════════════════════════════════════════════════════════════════════════════
# 4. Excel 보고서 작성
# ════════════════════════════════════════════════════════════════════════════

# ── 시트1: 전체 정산 내역 ────────────────────────────────────────────────────
def write_sheet1(wb, df: pd.DataFrame):
    ws = wb.create_sheet("1. 전체 정산 내역")
    sheet_title(ws, "전체 정산 내역 (세금 구분별 부가세 검증 결과)",
                "오차 행은 빨간색 하이라이트 / 올바른 금액 자동 기재")

    headers = [
        "정산확정일", "협력사명", "사업장", "부서명",
        "상품명", "세금구분", "단위", "수량",
        "공급가액(원)", "부가세(원)", "올바른청구액(원)", "실제청구액(원)",
        "오차금액(원)", "오차여부"
    ]
    SR = 4
    bg_h = {"과세": BG["header"], "면세": "2E75B6", "비과세": "375623"}

    for c, h in enumerate(headers, 1):
        hc(ws.cell(SR, c), h, wrap=True)
    ws.row_dimensions[SR].height = 32

    show = ["정산확정일","협력사명","사업장","부서명","상품명",
            "세금구분","단위","정산수량",
            "공급가액","부가세","올바른청구액","정산금액",
            "오차금액","오차여부"]

    for r_idx, (_, row) in enumerate(df.iterrows()):
        r   = SR + 1 + r_idx
        err = bool(row.get("오차여부", False))
        bg  = BG["error_l"] if err else (BG["alt"] if r % 2 == 0 else None)

        vals = [
            row.get("정산확정일", ""),
            row.get("협력사명", ""),
            row.get("사업장", ""),
            row.get("부서명", ""),
            str(row.get("상품명", ""))[:30],
            row.get("세금구분", ""),
            row.get("단위", ""),
            row.get("정산수량", 0),
        ]
        for c, v in enumerate(vals, 1):
            fmt = NUM if c == 8 else None
            al  = "left" if c in (2, 3, 4, 5) else "center"
            dc(ws.cell(r, c), v, fmt, bg=bg, align=al)

        # 금액 컬럼
        for c_idx, key in enumerate(
            ["공급가액", "부가세", "올바른청구액", "정산금액"], 9
        ):
            val = row.get(key, 0)
            dc(ws.cell(r, c_idx), val, NUM, bg=bg)

        # 오차금액
        diff = row.get("오차금액", 0)
        dc(ws.cell(r, 13), diff, NUM_S, bold=err,
           bg=bg, color=FG["red"] if err else "000000")

        # 오차여부
        label = "오차" if err else "정상"
        dc(ws.cell(r, 14), label, bold=err,
           bg=BG["error_l"] if err else BG["ok_l"],
           align="center",
           color=FG["red"] if err else FG["green"])

        # 오차 행: 올바른 금액 셀에 추가 강조
        if err:
            ws.cell(r, 11).fill = PatternFill("solid", start_color="FFEB9C")
            ws.cell(r, 11).font = Font(name=FONT, bold=True,
                                       color="9C5700", size=10)

    last = SR + len(df)
    total_row(ws, last + 1, 14)
    for c, col in [(8,"정산수량"),(9,"공급가액"),(10,"부가세"),
                   (11,"올바른청구액"),(12,"정산금액"),(13,"오차금액")]:
        cl = get_column_letter(c)
        ws.cell(last+1, c).value = f"=SUM({cl}{SR+1}:{cl}{last})"
        ws.cell(last+1, c).number_format = NUM

    col_ws = [(1,14),(2,22),(3,18),(4,14),(5,28),(6,10),(7,8),(8,8),
              (9,18),(10,14),(11,18),(12,18),(13,14),(14,10)]
    for c, w in col_ws: cw(ws, c, w)
    ws.freeze_panes = f"A{SR+1}"

    # 범례
    lr = last + 3
    ws.cell(lr, 1, "범례").font = Font(name=FONT, bold=True, size=9)
    for i, (lbl, bg_c) in enumerate([
        ("오차 행 (빨간 하이라이트)", BG["error_l"]),
        ("올바른 금액 셀 (노란 배경)", "FFEB9C"),
        ("정상 행", BG["ok_l"]),
    ], 0):
        rr = lr + 1 + i
        ws.cell(rr, 1).fill   = PatternFill("solid", start_color=bg_c)
        ws.cell(rr, 1).border = BD()
        ws.cell(rr, 2, lbl).font = Font(name=FONT, size=9)


# ── 시트2: 어음 발행 대상 목록 ───────────────────────────────────────────────
def write_sheet2(wb, bill_df: pd.DataFrame):
    ws = wb.create_sheet("2. 어음 발행 대상 목록")

    if len(bill_df) == 0:
        sheet_title(ws, "어음 발행 대상 목록")
        ws.cell(4, 1, "비과세 항목이 없어 어음 발행 대상이 없습니다.").font = Font(
            name=FONT, size=12, color=BG["ok"])
        return

    발행건 = (bill_df["판별결과"] == "어음 발행 대상").sum()
    오차건 = (bill_df["판별결과"] == "오차 확인 필요").sum()
    total  = bill_df["합산청구액"].sum()

    sheet_title(ws, "어음 발행 대상 목록",
                f"발행 대상 {발행건}건 | 오차 확인 필요 {오차건}건 | "
                f"합계 {total:,.0f}원  (금액 내림차순)")

    headers = ["주문번호", "협력사명", "거래처 규모", "사업장",
               "정산확정일", "품목 수", "합산 청구액(원)", "오차 건수", "판별 결과"]
    SR = 4
    for c, h in enumerate(headers, 1):
        hc(ws.cell(SR, c), h, bg=BG["bill"], wrap=True)
    ws.row_dimensions[SR].height = 30

    for r, row in enumerate(bill_df.itertuples(), SR+1):
        is_bill = row.판별결과 == "어음 발행 대상"
        bg = BG["bill_l"] if is_bill else BG["warn_l"]

        dc(ws.cell(r, 1), str(row.주문번호),  align="center", bg=bg)
        dc(ws.cell(r, 2), row.협력사명,       align="left",   bg=bg)
        dc(ws.cell(r, 3), row.거래처규모,      align="center", bg=bg)
        dc(ws.cell(r, 4), row.사업장,         align="left",   bg=bg)
        dc(ws.cell(r, 5), row.정산확정일,      align="center", bg=bg)
        dc(ws.cell(r, 6), row.품목수,         NUM,            bg=bg)
        dc(ws.cell(r, 7), row.합산청구액, NUM, bold=True,      bg=bg,
           color=FG["dark"] if is_bill else FG["warn"])
        dc(ws.cell(r, 8), row.오차건수,       NUM,            bg=bg,
           color=FG["red"] if row.오차건수 > 0 else "000000")
        판별색 = FG["green"] if is_bill else FG["warn"]
        dc(ws.cell(r, 9), row.판별결과, bold=True, bg=bg,
           align="center", color=판별색)

    last = SR + len(bill_df); tot = last + 1
    total_row(ws, tot, 9)
    for c in [6, 7, 8]:
        cl = get_column_letter(c)
        ws.cell(tot, c).value = f"=SUM({cl}{SR+1}:{cl}{last})"
        ws.cell(tot, c).number_format = NUM

    for c, w in [(1,16),(2,22),(3,12),(4,18),(5,14),(6,8),(7,22),(8,10),(9,16)]:
        cw(ws, c, w)

    # 판별 기준 안내
    lr = tot + 2
    ws.cell(lr, 1, "어음 발행 판별 기준").font = Font(
        name=FONT, bold=True, size=10, color=BG["bill"])
    criteria = [
        (f"조건1. 비과세 합산 청구액 >= {어음_최소금액:,}원"),
        (f"조건2. 거래처 규모 = {' 또는 '.join(sorted(어음_대상규모))}"),
        ("조건3. 해당 주문번호 내 정산 오차 없음"),
        ("※ 과세·면세는 금액 무관하게 어음 발행 제외"),
    ]
    for i, txt in enumerate(criteria):
        ws.cell(lr + 1 + i, 1, txt).font = Font(name=FONT, size=9, color="595959")


# ── 시트3: 오차 항목 목록 ────────────────────────────────────────────────────
def write_sheet3(wb, df: pd.DataFrame):
    ws = wb.create_sheet("3. 오차 항목 목록")
    오차 = df[df["오차여부"] == True].copy()

    if len(오차) == 0:
        sheet_title(ws, "오차 항목 목록")
        ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=10)
        c = ws.cell(4, 1, "✓ 오차 없음  —  모든 항목의 부가세가 정상입니다.")
        c.font = Font(name=FONT, bold=True, size=14, color=BG["ok"])
        c.fill = PatternFill("solid", start_color=BG["ok_l"])
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 40
        return

    sheet_title(ws, "오차 항목 목록",
                f"총 {len(오차)}건 | 오차 합계: {오차['오차금액'].abs().sum():,.0f}원  "
                f"— 오차 행은 빨간 하이라이트 + 올바른 금액 기재")

    headers = [
        "정산확정일", "협력사명", "상품명", "세금구분",
        "수량", "공급가액(원)", "부가세(원)",
        "올바른 청구액(원)", "실제 청구액(원)", "오차금액(원)"
    ]
    SR = 4
    for c, h in enumerate(headers, 1):
        hc(ws.cell(SR, c), h, bg=FG["red"], wrap=True)
    ws.row_dimensions[SR].height = 30

    for r, (_, row) in enumerate(오차.iterrows(), SR+1):
        for c, (key, fmt, al) in enumerate([
            ("정산확정일", None,  "center"),
            ("협력사명",   None,  "left"),
            ("상품명",     None,  "left"),
            ("세금구분",   None,  "center"),
            ("정산수량",   NUM,   "right"),
            ("공급가액",   NUM,   "right"),
            ("부가세",     NUM,   "right"),
            ("올바른청구액", NUM, "right"),
            ("정산금액",   NUM,   "right"),
            ("오차금액",   NUM_S, "right"),
        ], 1):
            val = row.get(key)
            if key == "상품명":
                val = str(val)[:30] if val else ""
            elif key == "올바른청구액":
                key = "올바른청구액"
            err_color = FG["red"] if key == "오차금액" else "000000"
            dc(ws.cell(r, c), val, fmt, bg=BG["error_l"],
               align=al, color=err_color)
        # 올바른 청구액 셀 강조
        ws.cell(r, 8).fill = PatternFill("solid", start_color="FFEB9C")
        ws.cell(r, 8).font = Font(name=FONT, bold=True, color="9C5700", size=10)

    last = SR + len(오차); tot = last + 1
    total_row(ws, tot, 10)
    for c in [5, 6, 7, 8, 9, 10]:
        cl = get_column_letter(c)
        ws.cell(tot, c).value = f"=SUM({cl}{SR+1}:{cl}{last})"
        ws.cell(tot, c).number_format = NUM

    for c, w in [(1,14),(2,22),(3,28),(4,10),(5,8),(6,18),(7,14),
                 (8,18),(9,18),(10,16)]:
        cw(ws, c, w)
    ws.freeze_panes = f"A{SR+1}"


# ── 시트4: 요약 ──────────────────────────────────────────────────────────────
def write_sheet4(wb, df: pd.DataFrame, bill_df: pd.DataFrame):
    ws = wb.create_sheet("4. 요약", 0)
    sheet_title(ws, "정산 요약 보고서",
                f"정산 기준월: {기준월}  |  파일 생성일: "
                f"{datetime.today().strftime('%Y년 %m월 %d일 %H:%M')}")
    ws.sheet_view.showGridLines = False

    # ── 세금 구분별 집계 ──────────────────────────────────────────────────
    SR = 4
    ws.cell(SR, 1, "[ 세금 구분별 집계 ]").font = Font(
        name=FONT, bold=True, size=12, color=BG["title"])
    ws.row_dimensions[SR].height = 24

    for c, h in enumerate(["세금 구분", "건수", "공급가액 합계(원)",
                            "부가세 합계(원)", "청구액 합계(원)",
                            "오차 건수", "오차 금액(원)"], 1):
        hc(ws.cell(SR+1, c), h, wrap=True)
    ws.row_dimensions[SR+1].height = 28

    세금_bg = {"과세": BG["alt"], "면세": "E2EFDA", "비과세": "EAF2FB",
               "미확인": BG["grey"]}
    grand_total = {"건수": 0, "공급가액": 0, "부가세": 0,
                   "청구액": 0, "오차건수": 0, "오차금액": 0}

    tax_row = SR + 2
    for 구분 in ["과세", "면세", "비과세"]:
        sub = df[df["세금구분"] == 구분]
        if len(sub) == 0:
            continue
        bg = 세금_bg[구분]
        cnt     = len(sub)
        공가합   = sub["공급가액"].sum()
        부가합   = sub["부가세"].sum()
        청구합   = sub["정산금액"].sum()
        오차_cnt = int(sub["오차여부"].sum())
        오차_amt = sub.loc[sub["오차여부"], "오차금액"].abs().sum()

        dc(ws.cell(tax_row, 1), 구분, bold=True, bg=bg, align="center")
        dc(ws.cell(tax_row, 2), cnt, NUM, bg=bg)
        dc(ws.cell(tax_row, 3), 공가합, NUM, bg=bg)
        dc(ws.cell(tax_row, 4), 부가합, NUM, bg=bg)
        dc(ws.cell(tax_row, 5), 청구합, NUM, bg=bg)
        err_bg = BG["error_l"] if 오차_cnt > 0 else bg
        dc(ws.cell(tax_row, 6), 오차_cnt, NUM, bg=err_bg,
           color=FG["red"] if 오차_cnt > 0 else "000000", bold=오차_cnt > 0)
        dc(ws.cell(tax_row, 7), 오차_amt, NUM, bg=err_bg,
           color=FG["red"] if 오차_amt > 0 else "000000", bold=오차_amt > 0)

        grand_total["건수"]   += cnt
        grand_total["공급가액"] += 공가합
        grand_total["부가세"]  += 부가합
        grand_total["청구액"]  += 청구합
        grand_total["오차건수"] += 오차_cnt
        grand_total["오차금액"] += 오차_amt
        tax_row += 1

    # 소계
    total_row(ws, tax_row, 7)
    for c, val in [(2, grand_total["건수"]),
                   (3, grand_total["공급가액"]),
                   (4, grand_total["부가세"]),
                   (5, grand_total["청구액"]),
                   (6, grand_total["오차건수"]),
                   (7, grand_total["오차금액"])]:
        ws.cell(tax_row, c).value  = val
        ws.cell(tax_row, c).number_format = NUM
        ws.cell(tax_row, c).font   = Font(name=FONT, bold=True)

    # 오차 건수 셀 색상 강조
    오차건수_셀 = ws.cell(tax_row, 6)
    오차건수_셀.fill = PatternFill("solid",
        start_color=BG["error_l"] if grand_total["오차건수"] > 0 else BG["ok_l"])
    오차건수_셀.font = Font(name=FONT, bold=True,
        color=FG["red"] if grand_total["오차건수"] > 0 else FG["green"])

    # ── 어음 발행 요약 ────────────────────────────────────────────────────
    bill_row = tax_row + 3
    ws.cell(bill_row, 1, "[ 어음 발행 대상 요약 ]").font = Font(
        name=FONT, bold=True, size=12, color=BG["bill"])
    ws.row_dimensions[bill_row].height = 24

    if len(bill_df) > 0:
        발행대상 = bill_df[bill_df["판별결과"] == "어음 발행 대상"]
        오차대상  = bill_df[bill_df["판별결과"] == "오차 확인 필요"]

        for r_off, (label, sub, bg, col) in enumerate([
            ("어음 발행 대상",    발행대상, BG["bill_l"], FG["dark"]),
            ("오차 확인 필요",    오차대상,  BG["warn_l"],  FG["warn"]),
        ], 1):
            r = bill_row + r_off
            dc(ws.cell(r, 1), label, bold=True, bg=bg, align="center",
               color=col)
            dc(ws.cell(r, 2), len(sub), NUM, bg=bg)
            dc(ws.cell(r, 3), sub["합산청구액"].sum(), NUM, bg=bg,
               bold=True, color=col)
            for c in range(4, 8): ws.cell(r, c).fill = PatternFill(
                "solid", start_color=bg)
    else:
        dc(ws.cell(bill_row+1, 1), "비과세 항목 없음 — 어음 발행 대상 없음",
           align="center", bg=BG["grey"])

    # ── 정산 기준 정보 ────────────────────────────────────────────────────
    info_row = (bill_row + 5 if len(bill_df) > 0 else bill_row + 3)
    ws.cell(info_row, 1, "[ 정산 기준 정보 ]").font = Font(
        name=FONT, bold=True, size=12, color=BG["title"])
    ws.row_dimensions[info_row].height = 24

    infos = [
        ("정산 기준월",     기준월),
        ("파일 생성일시",   datetime.today().strftime("%Y-%m-%d %H:%M:%S")),
        ("분석 대상 행수",  f"{len(df):,}행"),
        ("오차 허용 범위",  f"±{VAT_오차허용}원"),
        ("어음 발행 최소금액", f"{어음_최소금액:,}원"),
        ("어음 발행 대상 규모", " / ".join(sorted(어음_대상규모))),
    ]
    for i, (k, v) in enumerate(infos):
        r = info_row + 1 + i
        dc(ws.cell(r, 1), k, bold=True, bg=BG["alt"], align="left")
        dc(ws.cell(r, 2), v, align="left", bg=BG["grey"])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)

    for c, w in [(1,22),(2,18),(3,18),(4,18),(5,18),(6,14),(7,14)]:
        cw(ws, c, w)


# ════════════════════════════════════════════════════════════════════════════
# 5. 메인
# ════════════════════════════════════════════════════════════════════════════
def main():
    print("=" * 60)
    print("  KT 세금 정산 자동화 스크립트")
    print(f"  기준월: {기준월}")
    print("=" * 60)

    data      = load_data()
    validated = validate_vat(data)
    bill_df   = determine_bill(validated)

    wb = openpyxl.Workbook()
    del wb["Sheet"]

    write_sheet4(wb, validated, bill_df)          # 시트4 → index 0 (맨 앞)
    write_sheet1(wb, validated)
    write_sheet2(wb, bill_df)
    write_sheet3(wb, validated)

    wb.save(OUTPUT_FILE)

    오차수 = int(validated["오차여부"].sum())
    발행수 = int((bill_df["판별결과"] == "어음 발행 대상").sum()) if len(bill_df) > 0 else 0

    print(f"\n[저장 완료] {OUTPUT_FILE}")
    print("-" * 60)
    print(f"  1단계 부가세 오차  : {오차수}건")
    print(f"  2단계 어음 발행    : {발행수}건")
    print(f"  출력 시트          : 4개 (요약 / 전체내역 / 어음 / 오차)")
    print("=" * 60)

    if sys.platform == "win32":
        try:
            os.startfile(OUTPUT_FILE)
        except OSError:
            print(f"  파일 경로: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
