"""
발송 직전 점검(돈현님 요청, 2026-09-04, 5차): 애니콤에 함께 보낼 213종 표준BoM 작성양식
원본(input/7cb45e7e-Aviat_MDM_________BoM_____.xlsx)을 열어보니, 확인요청서 [6]에서 "정정해
달라"고 요청한 구코드 4건이 양식 자체에는 여전히 구코드로 남아 있었다(원본 그대로, 수정하지
않음 - 요청서 원칙 준수). 원본은 건드리지 않고, K코드 4건만 신코드로 바꾼 별도 사본을 만들어
발송용으로 제공한다.
입력: input/7cb45e7e-Aviat_MDM_________BoM_____.xlsx (읽기 전용)
출력: output/Aviat_표준BoM_작성양식(코드정정).xlsx
"""
from pathlib import Path

import openpyxl

BASE_DIR = Path("/home/user/kt-settlement-automation")
SRC = BASE_DIR / "input" / "7cb45e7e-Aviat_MDM_________BoM_____.xlsx"
OUT = BASE_DIR / "output" / "Aviat_표준BoM_작성양식(코드정정).xlsx"

# (순번, 구코드, 신코드, 품명) - 05_Aviat_공급사회신_검증.xlsx K코드_일치검증 시트 기준
FIXES = [
    (18, "K9090176", "K9210281", "HAX_커넥터 KIT_IF 케이블 종단용"),
    (21, "K9090179", "K9210280", "HAX_Arrestor KIT(N형 M,F콘넥터 각1EA 포함)"),
    (34, "K9093723", "K9210279", "HAX_Gigabit Ethernet SFP"),
    (80, "K9188156", "K9210278", "SFP TSoP STM-1/OC3 over Gig-E Module"),
]
HEADER_ROW = 4
SEQ_COL = 1
CODE_COL = 3
NAME_COL = 4


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb["표준BoM_213종"]

    for seq, old_code, new_code, expected_name in FIXES:
        row = seq + HEADER_ROW
        seq_cell = ws.cell(row=row, column=SEQ_COL).value
        code_cell = ws.cell(row=row, column=CODE_COL)
        name_cell = ws.cell(row=row, column=NAME_COL).value
        assert seq_cell == seq, f"순번 불일치: row={row} 실제={seq_cell} 기대={seq}"
        assert code_cell.value == old_code, f"K코드 불일치: row={row} 실제={code_cell.value} 기대={old_code}"
        assert name_cell == expected_name, f"품명 불일치: row={row} 실제={name_cell} 기대={expected_name}"
        code_cell.value = new_code
        print(f"순번 {seq}: {old_code} -> {new_code} ({expected_name})")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    # 저장 후 재오픈해 값이 실제로 반영됐는지, 원본은 그대로인지 확인
    check = openpyxl.load_workbook(OUT, data_only=True)["표준BoM_213종"]
    orig = openpyxl.load_workbook(SRC, data_only=True)["표준BoM_213종"]
    for seq, old_code, new_code, _ in FIXES:
        row = seq + HEADER_ROW
        assert check.cell(row=row, column=CODE_COL).value == new_code, "정정 반영 실패"
        assert orig.cell(row=row, column=CODE_COL).value == old_code, "원본이 변경됨(금지 원칙 위반)"

    print(f"\n=== K코드 4건 정정 완료, 원본은 변경하지 않음 ===")
    print(f"원본(변경 없음): {SRC}")
    print(f"정정본(발송용): {OUT}")


if __name__ == "__main__":
    main()
