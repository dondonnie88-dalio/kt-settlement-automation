"""
8단계: Aviat와 Ceragon 구성방식별 가격 비교
입력: output/04_Aviat_구성방식별_표준BoM_추정.xlsx, output/06_Ceragon_구성방식별_BoM_정리.xlsx
출력: output/07_Aviat_Ceragon_구성가격비교.xlsx

핵심 제약(요약 시트에도 기록):
- 돈현님 피드백(2026-09-04) 반영: 8GHz/11GHz는 Aviat·Ceragon 양측 자료에 모두 존재하므로
  "대역이 겹치지 않는다"고 쓰지 않는다. 04v2(8GHz/11GHz IAP3 재분석)에서 Aviat 측 2+0/4+0
  후보 근거(유력/참고 후보)를 확인했으나 완성된 표준 BoM 수준은 아니어서, '대역은 겹치지만
  Aviat 완성 BoM 미확정으로 구성 총액 비교가 보류됨'으로 정리한다. 반면 6+0/8+0(8/11GHz)은
  그 배수 자체를 시사하는 Aviat 발주이력이 전혀 없어 비교 불가이고, Aviat의 4GHz(WBX/CTR
  텍스트 앵커) 전 구성은 Ceragon 쪽에 4GHz 수량 자료 자체가 없어 비교 불가다. 이 세 가지는
  서로 다른 사유이므로 시트별로 구분해 기록한다(임의로 다른 대역을 묶어 비교하지 않는다).
- Ceragon 파일에는 단가(공급단가)만 있고 별도의 '매입가'가 없다. 과제 지침에 따라 이 값은
  '판매가 또는 계약단가'로만 사용하고, 매입가 비교에는 사용하지 않는다
  ('Ceragon 매입가 미제공'으로 표시).
- 개별 품목 중 기능·설명이 명확히 동일한 것만 '직접대응_품목'으로 비교한다.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_style import finalize_sheet, ERROR_FILL, REVIEW_FILL, CONFIRMED_FILL, FMT_AMOUNT, mark_fill

BASE_DIR = Path("/home/user/kt-settlement-automation")
OUTPUT_DIR = BASE_DIR / "output"
SRC_AVIAT_BOM = OUTPUT_DIR / "04_Aviat_구성방식별_표준BoM_추정.xlsx"
SRC_AVIAT_BOM_V2 = OUTPUT_DIR / "04_Aviat_구성방식별_표준BoM_추정_v2.xlsx"
SRC_CERAGON_BOM = OUTPUT_DIR / "06_Ceragon_구성방식별_BoM_정리.xlsx"

CONFIGS = ["2+0", "4+0", "6+0", "8+0"]

# 근거가 명확한 개별 품목 직접 대응 쌍(K코드는 Aviat / Ceragon 각각의 원본 K코드)
DIRECT_PAIRS = [
    ("K9198328", "K9183522", "옥외(Outdoor)용 산업용 1G GbE SFP 광모듈",
     "양측 모두 '옥외 Outdoor용 Optic Gbit Module' 설명이 동일. 온도사양은 Aviat -40~+85℃, "
     "Ceragon -30~+50℃로 약간 다름(확인 필요)."),
    ("K9210279", "K9178770", "실내용 1G GbE SFP 광모듈",
     "양측 모두 실내형 1G Gigabit Ethernet SFP. 거리/파장 사양은 Aviat 자료에 명시되어 있지 않음"
     "(확인 필요)."),
    ("K9198347", "K9197307", "10G SFP+ 광모듈(참고용, 완전 동일 사양 아님)",
     "Aviat 품목은 산업용/실외 등급(-40~+85℃)이나 Ceragon 품목은 실외 등급 명시가 없어 "
     "완전히 동일한 사양은 아님 - 참고용 비교(확인 필요)."),
]


def load_aviat_config_totals():
    wb = openpyxl.load_workbook(SRC_AVIAT_BOM, data_only=True)
    result = {}
    for cfg in CONFIGS:
        ws = wb[f"{cfg}_BoM"]
        headers = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(headers)}
        items = []
        sale_total = 0.0
        buy_total = 0.0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[idx["K코드"]]:
                continue
            qty = row[idx["1개 링크 기준 추정 수량(대표값)"]]
            qty_num = qty if isinstance(qty, (int, float)) else (1 if isinstance(qty, str) and "앵커" in qty else None)
            sale_amt = row[idx["판매금액"]] or 0
            buy_amt = row[idx["매입금액"]] or 0
            items.append({
                "K코드": row[idx["K코드"]], "품명": row[idx["품명"]],
                "수량": qty, "확정수준": row[idx["확정 수준"]],
                "판매금액": sale_amt, "매입금액": buy_amt,
            })
            if isinstance(sale_amt, (int, float)):
                sale_total += sale_amt
            if isinstance(buy_amt, (int, float)):
                buy_total += buy_amt
        result[cfg] = {"items": items, "판매총액": sale_total, "매입총액": buy_total,
                        "주파수": "4GHz(WBX 근거) 또는 미상(계약서 텍스트 근거만 있음, 발주이력 없음)"}
    return result


def load_aviat_iap3_summary():
    """04v2(8GHz/11GHz IAP3 재분석) 결과를 요약해, 04(v1)이 놓친 8/11GHz 후보 근거를
    구성총액 비교 사유에 반영한다. 04v2는 2+0/4+0만 다루므로(6+0/8+0은 발주이력 근거 없음),
    이 두 구성만 값을 채우고 나머지는 None으로 남긴다."""
    wb = openpyxl.load_workbook(SRC_AVIAT_BOM_V2, data_only=True)
    result = {}
    for band in ["8GHz", "11GHz"]:
        for cfg in ["2+0", "4+0"]:
            ws = wb[f"{band}_IAP3_{cfg}"]
            headers = [c.value for c in ws[1]]
            idx = {h: i for i, h in enumerate(headers)}
            n_total = n_strong = n_ref = 0
            strong_sale_total = 0.0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[idx["K코드"]]:
                    continue
                n_total += 1
                level = row[idx["판정수준"]]
                if level == "유력 후보":
                    n_strong += 1
                    # 주의(2026-09-08 정정): '구성비 배수' 열은 N당 배수(예: 1=채널당 1개)이지
                    # 링크 전체 수량이 아니다. 이걸 그대로 단가에 곱하면 2+0은 실제의 1/2, 4+0은
                    # 1/4로 총액이 축소된다 - 반드시 '실제 수량(N+0 링크 전체)' 열을 써야 한다.
                    qty = row[idx["실제 수량(N+0 링크 전체, 양쪽 사이트 합산)"]]
                    price = row[idx["판매단가"]] or 0
                    if isinstance(qty, (int, float)):
                        strong_sale_total += qty * price
                elif level == "참고 후보":
                    n_ref += 1
            result[(band, cfg)] = {
                "품목수": n_total, "유력후보": n_strong, "참고후보": n_ref,
                "유력후보_참고총액": strong_sale_total if n_strong > 0 else None,
            }
    return result


def load_ceragon_config_totals():
    wb = openpyxl.load_workbook(SRC_CERAGON_BOM, data_only=True)
    ws = wb["구성별_총액"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[idx["원본주파수"]]:
            continue
        key = (row[idx["원본주파수"]], row[idx["표준화구성표기"]], row[idx["SD구분"]])
        result[key] = {
            "품목종류수": row[idx["품목종류수"]], "총수량": row[idx["총수량"]],
            "기존금액합계": row[idx["기존금액합계"]], "인하금액합계": row[idx["인하금액합계"]],
        }
    return result


def load_price_lookup(path, sheet, kcode_col, name_col, sale_col, buy_col=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    lut = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        k = row[idx[kcode_col]] if kcode_col in idx else None
        if not k:
            continue
        lut[k] = {
            "품명": row[idx[name_col]] if name_col in idx else None,
            "판매가": row[idx[sale_col]] if sale_col in idx else None,
            "매입가": row[idx[buy_col]] if buy_col and buy_col in idx else None,
        }
    return lut


def main():
    aviat_cfg = load_aviat_config_totals()
    aviat_iap3 = load_aviat_iap3_summary()
    ceragon_cfg = load_ceragon_config_totals()
    aviat_price = load_price_lookup(OUTPUT_DIR / "01_Aviat_213종_정규화.xlsx", "원본정리",
                                     "K코드", "품명", "판매단가", "매입단가")
    ceragon_price = load_price_lookup(SRC_CERAGON_BOM, "품목마스터", "K코드", "품명",
                                       "최종인하단가", None)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------------- 구성총액_비교 ----------------
    ws1 = wb.create_sheet("구성총액_비교")
    headers1 = ["구성방식", "Aviat 대역/근거", "Aviat 판매총액(추정)", "Aviat 매입총액(추정)",
                "Ceragon 대역", "Ceragon SD구분", "Ceragon 판매총액(단가기준)", "Ceragon 매입총액",
                "직접 비교 가능 여부", "사유"]
    ws1.append(headers1)
    for cfg in CONFIGS:
        for band in ["8GHz", "11GHz"]:
            for sd in ["SD", "Non_SD"]:
                cer = ceragon_cfg.get((band, cfg, sd))
                r = ws1.max_row + 1
                iap3 = aviat_iap3.get((band, cfg))
                if iap3 is None:
                    # 6+0/8+0: 대역(8/11GHz) 자체는 Aviat에도 존재하나, 이 배수를 시사하는
                    # 발주이력 근거가 전혀 없다(04v2 확인필요 시트 참조) - "대역이 다르다"가
                    # 아니라 "이 배수의 구성표기·수량 근거 자체가 없다"가 정확한 사유다.
                    aviat_band_note = f"{band} 자체는 Aviat 자료에도 존재하나, {cfg} 배수를 " \
                                       "시사하는 발주이력 근거가 확인되지 않음(04v2 확인필요 시트 참조)"
                    aviat_sale = None
                    comparable = "아니오(근거 없음)"
                    reason = (
                        f"대역({band})은 Aviat·Ceragon 양측 자료에 모두 존재하므로 '대역이 겹치지 "
                        f"않는다'고 표현하지 않는다. 다만 {cfg} 구성 자체를 시사하는 Aviat 발주이력이 "
                        f"없어(04v2 재분석 결과 포함) 이 조합은 비교 대상이 될 수 없음."
                    )
                elif iap3["유력후보"] > 0:
                    aviat_band_note = (f"{band} {cfg}: 04v2 재분석 결과 유력후보 {iap3['유력후보']}건/"
                                        f"참고후보 {iap3['참고후보']}건(표준 BoM 확정 아님)")
                    aviat_sale = iap3["유력후보_참고총액"]
                    comparable = "보류(대역 겹침, Aviat 미확정)"
                    reason = (
                        f"대역({band})은 Aviat·Ceragon 양측에 모두 존재함(대역 자체는 겹침). 04v2 "
                        f"재분석에서 {cfg} 구성의 핵심 품목 {iap3['유력후보']}건이 '유력 후보' 수준까지 "
                        f"확인되었으나(반복 관측 비율 안정), 전체 품목({iap3['품목수']}건) 중 다수는 "
                        f"아직 참고 후보 단계이고 완성된 표준 BoM으로 확정되지 않아, 구성 총액 단위의 "
                        f"전면 비교는 보류한다. 왼쪽 '판매총액' 값은 유력후보 품목만의 부분 합계(참고용)"
                        f"이며 Ceragon 쪽 전체 품목 총액과 직접 비교 가능한 값이 아니다."
                    )
                else:
                    aviat_band_note = (f"{band} {cfg}: 04v2 재분석 결과 참고후보 {iap3['참고후보']}건만 "
                                        "확인(반복 관측으로 뒷받침되는 유력 후보 없음)")
                    aviat_sale = None
                    comparable = "보류(대역 겹침, Aviat 근거 부족)"
                    reason = (
                        f"대역({band})은 Aviat·Ceragon 양측에 모두 존재함(대역 자체는 겹침). 다만 "
                        f"04v2 재분석에서 {cfg} 구성 품목 대부분이 다중대역 혼재 윈도우 등으로 인해 "
                        f"'참고 후보' 수준에 그쳐, 총액을 산출할 만큼의 근거가 아직 없음. 임의로 수량을 "
                        f"채워 총액을 계산하지 않는다(공급사 확인 필요)."
                    )
                ws1.append([
                    cfg, aviat_band_note, aviat_sale, None,
                    band, sd, cer["인하금액합계"] if cer else None, "Ceragon 매입가 미제공",
                    comparable, reason,
                ])
                mark_fill(ws1, r, 9, REVIEW_FILL if iap3 is not None else ERROR_FILL)
    for c in [3, 4, 7]:
        for r in range(2, ws1.max_row + 1):
            ws1.cell(row=r, column=c).number_format = FMT_AMOUNT
    finalize_sheet(ws1, 1, len(headers1), ws1.max_row)

    # ---------------- 품목구성_차이 ----------------
    ws2 = wb.create_sheet("품목구성_차이")
    ws2.append(["항목", "Aviat", "Ceragon", "근거/비고"])
    diffs = [
        ("구성 확장 방식", "모듈(카드) 추가형 - IDU 셀프에 모뎀/인터페이스 카드를 추가해 채널 증설",
         "시스템(셀프) 단위 - 1RU(Non-SD)/2RU(SD) 셀프 자체를 교체하며 채널 증설",
         "공급사 미팅 녹취(⑦ 파일) P17, P60 참조. 구두 설명으로 확인 필요 등급."),
        ("SD 셀프 구분", "SD/Non-SD를 구분하는 별도 셀프 코드가 확정되어 있지 않음(확인 필요)",
         "1RU=Non-SD 전용, 2RU=SD 전용으로 원본 파일에 명시(확정)",
         "60d6ff1e 파일 header row6 'IP-20N 1RU/2RU적용' 표기."),
        ("6+0 구성 근거", "계약서/발주이력 어디에도 6+0을 지시하는 표기 없음(확인 필요)",
         "6GHz 대역 및 8GHz_11GHz 대역 모두 'Configuration SW Package' 라이선스가 6+0까지 "
         "명시적으로 존재(확정, 단 하드웨어 수량 매트릭스는 6GHz 대역에서 공란)",
         "60d6ff1e 파일 8G_가격인하 시트 row488~489 등."),
        ("주파수 커버리지", "6/8/11GHz(L6/U6 세분) + 4GHz(WTM4500XT)까지 폭넓게 계약되어 있음",
         "제공 자료 기준 8GHz/11GHz 수량 매트릭스만 존재, 4GHz 자체는 확인 안 됨(확인 필요)",
         "1단계 구조조사 참조."),
        ("가격 구조(공급사 발언)", "-",
         "저구성(1+0/2+0)은 원가 근접/이하로 책정하고 설치자재(Waveguide/Coupler 등)에서 "
         "가격을 보정한다고 발언 - 품목 단위 개별비교 시 왜곡 가능성 있음",
         "공급사 미팅 녹취 P45-46, P132-142 (구두 설명, 확인 필요)."),
    ]
    for row in diffs:
        ws2.append(list(row))
    finalize_sheet(ws2, 1, 4, ws2.max_row)
    ws2.column_dimensions["B"].width = 45
    ws2.column_dimensions["C"].width = 45
    ws2.column_dimensions["D"].width = 45

    # ---------------- 직접대응_품목 ----------------
    ws3 = wb.create_sheet("직접대응_품목")
    headers3 = ["비교대상", "Aviat K코드", "Aviat 품명", "Aviat 판매가", "Aviat 매입가",
                "Ceragon K코드", "Ceragon 품명", "Ceragon 단가(판매가/계약단가)",
                "판매가 차액", "판매가 차이율", "비고"]
    ws3.append(headers3)
    for aviat_k, ceragon_k, label, note in DIRECT_PAIRS:
        av = aviat_price.get(aviat_k, {})
        cer = ceragon_price.get(ceragon_k, {})
        r = ws3.max_row + 1
        ws3.append([
            label, aviat_k, av.get("품명"), av.get("판매가"), av.get("매입가"),
            ceragon_k, cer.get("품명"), cer.get("판매가"), None, None, note,
        ])
        ws3.cell(row=r, column=9).value = f"=D{r}-H{r}"
        ws3.cell(row=r, column=10).value = f"=IF(D{r}=0,0,I{r}/D{r})"
    for c in [4, 5, 8, 9]:
        for r in range(2, ws3.max_row + 1):
            ws3.cell(row=r, column=c).number_format = FMT_AMOUNT
    for r in range(2, ws3.max_row + 1):
        ws3.cell(row=r, column=10).number_format = "0.0%"
    finalize_sheet(ws3, 1, len(headers3), ws3.max_row)

    # ---------------- 판매가_비교 ----------------
    ws4 = wb.create_sheet("판매가_비교")
    ws4.append(["비교 유형", "내용", "Aviat", "Ceragon(판매가/계약단가로 사용)", "비고"])
    ws4.append(["개별품목(직접대응)", "SFP 등 직접대응 3건 평균 판매가 차이", "직접대응_품목 시트 참조",
                "직접대응_품목 시트 참조", "3건 모두 Ceragon 단가가 낮게 나타남(확인 필요, 정식 견적 아님)"])
    ws4.append(["구성총액(2+0~8+0)", "동일 대역·동일 구성 기준 총액 비교",
                "8GHz/11GHz 대역 자체는 존재(04v2 IAP3 재분석). 2+0/4+0은 유력/참고 후보 수준 "
                "근거 있으나 완성 BoM 미확정, 6+0/8+0은 근거 자체 없음",
                "8GHz/11GHz 매트릭스 존재(2+0~8+0 전 구성)",
                "구성총액_비교 시트: 2+0/4+0(8·11GHz)은 '보류'(대역은 겹침, Aviat 미확정), "
                "6+0/8+0(8·11GHz) 및 4GHz 전 구성(WBX/CTR 앵커)은 '비교불가' - 자세한 사유는 "
                "구성총액_비교/비교불가 시트 참조"])
    finalize_sheet(ws4, 1, 5, ws4.max_row)
    ws4.column_dimensions["B"].width = 40
    ws4.column_dimensions["E"].width = 50

    # ---------------- 매입가_비교 ----------------
    ws5 = wb.create_sheet("매입가_비교")
    ws5.append(["비교 유형", "내용", "Aviat 매입가", "Ceragon 매입가", "비고"])
    for aviat_k, ceragon_k, label, note in DIRECT_PAIRS:
        av = aviat_price.get(aviat_k, {})
        ws5.append([label, f"{aviat_k} vs {ceragon_k}", av.get("매입가"), "Ceragon 매입가 미제공",
                    "Ceragon 파일에는 공급단가만 있고 KT 매입가가 별도로 제공되지 않음"])
        mark_fill(ws5, ws5.max_row, 4, REVIEW_FILL)
    for cfg in CONFIGS:
        ws5.append([f"{cfg} 구성총액", "-", aviat_cfg[cfg]["매입총액"] or None, "Ceragon 매입가 미제공",
                    "동일 사유"])
        mark_fill(ws5, ws5.max_row, 4, REVIEW_FILL)
    for r in range(2, ws5.max_row + 1):
        ws5.cell(row=r, column=3).number_format = FMT_AMOUNT
    finalize_sheet(ws5, 1, 5, ws5.max_row)
    ws5.column_dimensions["B"].width = 30
    ws5.column_dimensions["E"].width = 50

    # ---------------- 비교불가 ----------------
    # 주의: 8GHz/11GHz는 Aviat·Ceragon 양측 자료에 모두 존재하므로(04v2 참조) "대역이 다르다"는
    # 이 시트에 다시 쓰지 않는다. 2+0/4+0은 04v2에서 후보 근거가 있어 '불가'가 아니라 '보류'
    # 상태이므로(구성총액_비교 시트 참조) 이 시트에는 올리지 않는다. 여기 남는 것은 (a) 6+0/8+0처럼
    # 그 배수 자체를 시사하는 Aviat 발주이력이 전혀 없는 경우, (b) Ceragon 쪽에 4GHz 데이터 자체가
    # 없어 Aviat의 4GHz(WBX/CTR 텍스트 앵커) 자료와 맞대볼 상대가 없는 경우, 두 가지뿐이다.
    ws6 = wb.create_sheet("비교불가")
    headers6 = ["구성방식", "대역 조합", "비교불가 사유"]
    ws6.append(headers6)
    for cfg in CONFIGS:
        for band in ["8GHz", "11GHz"]:
            if aviat_iap3.get((band, cfg)) is not None:
                continue
            ws6.append([cfg, f"Aviat({band}) vs Ceragon({band})",
                        f"{band} 대역 자체는 Aviat 자료에도 존재하나(04v2 참조), {cfg} 배수를 시사하는 "
                        "Aviat 발주이력 근거가 전혀 확인되지 않음(대역 불일치가 아니라 이 배수의 근거 "
                        "부족 - 공급사 확인 필요)."])
        ws6.append([cfg, "Aviat(WBX/CTR 텍스트 앵커, 4GHz 위주) vs Ceragon(4GHz)",
                    "Ceragon 제공 자료에 4GHz 대역 수량 매트릭스 자체가 없음(라이선스 단가 일부 제외). "
                    "이는 Ceragon 쪽 4GHz 데이터 부재가 원인이며, Aviat·Ceragon 대역이 전반적으로 "
                    "겹치지 않는다는 뜻은 아니다(8GHz/11GHz는 양측에 모두 존재 - 구성총액_비교 시트 참조)."])
    finalize_sheet(ws6, 1, len(headers6), ws6.max_row)
    ws6.column_dimensions["C"].width = 70

    # ---------------- 공급사확인사항 ----------------
    # 돈현님 피드백(2026-09-04, 2차): 213종 전체 수량 회신보다 아래 5개 질문이 더 중요함 -
    # 이 5개를 최우선으로 배치하고, 근거가 되는 04v2 관측치를 각 질문에 구체적으로 인용한다.
    ws7 = wb.create_sheet("공급사확인사항")
    ws7.append(["번호", "확인 요청 내용"])
    asks = [
        "[1순위] Aviat: 8GHz IAP3 기준 2+0/4+0 대표 표준 구성(수량 포함)은 무엇인가? - 04v2에서 "
        "8GHz 2+0은 EVT-003(2025-06-27) 1건, 8GHz 4+0은 EVT-013(2026-04-27) 1건에서 VR4 CHASSIS/"
        "ODU LOW·HIGH/FAN-CV/License/Mounting Bracket/Hybrid가 각각 정확히 동일 비율로 함께 "
        "관측됨(관측횟수 1건, 아직 독립 반복 없음) - 이 조합이 맞는지 확인 요청.",
        "[2순위] Aviat: 11GHz IAP3 기준 2+0/4+0 대표 표준 구성(수량 포함)은 무엇인가? - 04v2에서 "
        "11GHz 2+0은 EVT-007·EVT-010 두 독립 이벤트에서 VR4 CHASSIS/ODU LOW·HIGH/FAN-CV/Node "
        "License/Mounting Bracket 6개 품목이 반복 비율 1.00으로 확인되어 '유력 후보'로 분류함(가장 "
        "근거가 뚜렷함) - 이 조합의 정확성 확인 요청. 11GHz 4+0은 EVT-008 1건뿐이고 품목별 비율이 "
        "엇갈려(Chassis/Fan/License는 0.5, ODU/Modem은 1.0) 1개의 4채널 링크인지 2개의 2채널 링크가 "
        "섞인 것인지 불명확 - 확인 필요.",
        "[3순위] Aviat: VR4 적용 조건과 VR10 적용 조건은 무엇인가? - 두 계열이 언제 각각 선택되는지 "
        "(용량/거리/사이트 규모 등 기준) 발주이력만으로는 판단할 수 없음.",
        "[4순위] Aviat: WBX와 IAP3(VR4/VR10) 계열의 관계는 무엇인가? - 서로 대체 가능한 다른 세대 "
        "제품군인지, 용도가 구분된 별도 제품군인지, 필수/선택 여부 확인 요청.",
        "[5순위] Aviat: 6+0/8+0 구성이 실제로 공급 가능한가? - 계약품목·발주이력 어디에서도 6+0 "
        "표기 자체가 없고(2+0), 8+0은 IAP3 계열(8/11GHz)에서 발주 근거가 전혀 없음(WBX 계열 4GHz "
        "8채널 표기만 확정) - 실제 공급 가능 여부와 가능하다면 표준 BoM 확인 요청.",
        "Aviat: 213종 계약품목 전체에 대한 2+0/4+0/6+0/8+0 수량 기재 표준BoM 양식 회신(위 5개 "
        "질문보다 우선순위는 낮으나 여전히 필요).",
        "Ceragon: 4GHz 대역 지원 여부 및 지원 시 구성별 BoM/단가 제공.",
        "Ceragon: '매입가'(KT 조달원가) 별도 제공 여부 - 현재 자료는 단가(공급가)만 존재.",
        "Ceragon: 6GHz 대역 'Configuration SW Package' 라이선스가 실제 하드웨어 구성과 어떻게 "
        "연결되는지(현재 수량 매트릭스 공란) 확인.",
        "Ceragon: 6GHz 'Configuration SW Package' 8개 품목(신규코드-R484~R491) 설명에 적힌 "
        "'Nnd-Activation' 표기가 2+0 SD=2nd부터 8+0 Non_SD=9nd까지 순차 증가하는데, 같은 라이선스 "
        "계열인 8GHz/11GHz용 8개 품목(R492~R499)은 전부 '2nd-Activation'으로 동일하다 - 6GHz "
        "쪽만 값이 순차 증가하는 것은 원본 가격표 작성 시 행 복사 과정에서 번호가 밀린 것으로 "
        "보인다(2nd로 통일 여부 확인 필요, 6GHz 쪽 전체를 '2nd-Activation'으로 통일해야 하는지 "
        "확인 요청).",
        "직접대응_품목 3건(SFP류)의 정확한 온도/거리 사양 일치 여부 재확인.",
    ]
    for i, a in enumerate(asks, start=1):
        ws7.append([i, a])
    finalize_sheet(ws7, 1, 2, ws7.max_row)
    ws7.column_dimensions["B"].width = 100

    # ---------------- 요약 ----------------
    ws8 = wb.create_sheet("요약")
    ws8.append(["항목", "내용"])
    ws8.append(["현재 단계(돈현님 피드백 반영, 2026-09-04)", "이 파일은 아직 '가격 비교 자료'가 "
                                                       "아니라 'Aviat에 무엇을 확인해야 하는지 정리한 "
                                                       "자료' 단계다. 8GHz/11GHz IAP3 구성 후보는 "
                                                       "04v2에서 식별했지만 전부 유력/참고 후보이고 "
                                                       "'표준 BoM 확정' 품목은 아직 0건이므로, 가격 "
                                                       "총액 비교에 앞서 공급사확인사항 시트의 5개 "
                                                       "우선 질문에 대한 회신이 먼저 필요하다."])
    ws8.append(["비교 가능 범위", "Aviat와 Ceragon은 8GHz 및 11GHz 대역이 공통으로 존재한다. 다만 "
                              "Aviat의 구성방식별 완성 BoM이 아직 공식적으로 확인되지 않아 구성 총액 "
                              "비교는 보류한다. 개별 품목 단위 3건(SFP류)은 직접 비교 가능. 세부적으로는: "
                              "(1) 8GHz/11GHz 2+0/4+0 후보는 04v2에서 유력/참고 후보 수준까지만 확인되고 "
                              "완성된 표준 BoM으로 확정되지 않아, 구성 총액 단위의 전면 비교는 보류함. "
                              "(2) 6+0/8+0은 그 배수를 시사하는 Aviat 발주이력 근거 자체가 없어 비교 "
                              "불가. (3) Aviat의 4GHz(WBX/CTR 텍스트 앵커) 전 구성은 Ceragon 쪽에 4GHz "
                              "수량 자료 자체가 없어 비교 불가(자세한 내용은 구성총액_비교/비교불가 시트 "
                              "참조)."])
    ws8.append(["직접대응 3건 결과", "SFP류 3건 모두 Ceragon 단가가 Aviat 판매가보다 낮게 나타남(구체 "
                                "수치는 직접대응_품목 시트 참조). 단, 이는 공식 견적이 아니라 참고용 "
                                "단가 비교이므로 협상 근거로 바로 사용하지 말 것."])
    ws8.append(["가격축 원칙 준수", "Ceragon 매입가가 없어 Aviat 매입가는 어떤 Ceragon 수치와도 비교하지 "
                                "않았음('Ceragon 매입가 미제공'으로 명시). 판매가 비교만 수행."])
    ws8.append(["핵심 결론", "협상에 즉시 활용 가능한 것은 개별 SFP 단가 비교뿐이다. 지금 병목은 "
                          "가격이 아니라 BoM 확정이다: 11GHz 2+0은 독립된 2개 이벤트에서 반복 확인된 "
                          "'유력 후보'(VR4 CHASSIS/ODU LOW·HIGH/FAN-CV/Mounting Bracket/Node "
                          "License)까지 나왔고, 8GHz 2+0(EVT-003)과 8GHz 4+0(EVT-013)도 단일 "
                          "이벤트지만 동일하게 깨끗한 패턴을 보였다(공급사확인사항 1·2순위 질문 "
                          "참조). 다음 단계는 가격 비교가 아니라 이 3개 후보와 11GHz 4+0(EVT-008, "
                          "품목별 비율이 엇갈려 확인 필요)을 Aviat에 제시해 표준 BoM으로 확정받는 "
                          "것이다(다른 확인필요 사항은 5·6단계 결과 및 이 파일 공급사확인사항 시트 "
                          "참조)."])
    finalize_sheet(ws8, 1, 2, ws8.max_row)
    ws8.column_dimensions["B"].width = 100

    out_path = OUTPUT_DIR / "07_Aviat_Ceragon_구성가격비교.xlsx"
    wb.save(out_path)

    n_pending = sum(1 for row in ws1.iter_rows(min_row=2, values_only=True) if "보류" in row[8])
    n_blocked = sum(1 for row in ws1.iter_rows(min_row=2, values_only=True) if "아니오" in row[8])

    print("=== 8단계: Aviat-Ceragon 구성가격 비교 완료 ===")
    print(f"입력 파일: {SRC_AVIAT_BOM.name}, {SRC_AVIAT_BOM_V2.name}, {SRC_CERAGON_BOM.name}")
    print(f"처리: 구성 {len(CONFIGS)}종 × 대역 2 × SD구분 2 = {ws1.max_row-1}개 조합 검토, 직접대응 품목 {len(DIRECT_PAIRS)}건")
    print(f"생성 파일: {out_path}")
    print("주요 검증 결과:")
    print(f"  - 구성총액 비교(8/11GHz {ws1.max_row-1}개 조합): 확정 비교 0건 / 보류(대역 겹침, "
          f"Aviat 미확정) {n_pending}건 / 비교불가(근거 없음) {n_blocked}건 - '대역 불일치'가 아님")
    print(f"  - 4GHz(WBX/CTR 앵커) 전 구성: Ceragon에 4GHz 데이터가 없어 비교불가(비교불가 시트 참조)")
    print(f"  - 판매가 직접비교 가능 품목: {len(DIRECT_PAIRS)}건")
    print(f"  - 매입가 비교 가능 품목: 0건 (Ceragon 매입가 미제공)")
    print(f"  - 공급사 확인 필요 항목: {len(asks)}건")


if __name__ == "__main__":
    main()
