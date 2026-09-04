"""
3단계: Aviat 발주이력 정제
입력: input/3b46dc94-20252026_____.xlsx (시트 '주문현황', row6~250, 245건)
출력: output/02_Aviat_발주이력_정제.xlsx
시트: 정제데이터 / 제외검토 / 금액오류 / 코드매핑오류 / 주문요약

원본 파일은 읽기 전용으로만 사용한다. 개인정보(이메일/연락처/핸드폰)는 결과 파일에서 제외한다.
"""
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_style import (
    finalize_sheet, ORIGIN_FONT, ERROR_FILL, REVIEW_FILL, CONFIRMED_FILL,
    FMT_AMOUNT, FMT_DATE, mark_fill, mark_font,
)

BASE_DIR = Path("/home/user/kt-settlement-automation")
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
SRC = INPUT_DIR / "3b46dc94-20252026_____.xlsx"
AVIAT213_SRC = INPUT_DIR / "510331d5-_____MDMmini_digital_microwave_213______.xlsx"

DATA_START, DATA_END = 6, 250  # 245건

# 주문현황 열 인덱스 (1-based, 헤더 row5 기준)
C = {
    "주문번호": 1, "품목번호": 2, "IP/DIP": 3, "주문일자": 4, "주문상태": 5, "반품/교환": 6,
    "부서": 11, "주문자": 13, "상품코드": 22, "K코드": 23, "상품명": 24, "대표규격": 25,
    "모델명": 26, "제조사": 27, "단위": 30, "수량": 31, "판매단가": 32, "판매금액": 33,
    "매입단가": 34, "매입금액": 35, "매출구분": 37, "배송지": 41, "배송지상세": 42,
    "배송희망일": 53, "발주일자": 73, "취소수량": 79, "반품수량": 81, "입고수량": 86,
    "입고일자": 92, "빌딩명": 109, "국사명": 112, "프로젝트명": 116,
}


def parse_date(s):
    if s is None or str(s).strip() == "":
        return None, "공백"
    s = str(s).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    return None, f"형식불명({s})"


def parse_number(v):
    if v is None or str(v).strip() == "":
        return None, "공백"
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip().replace(",", "")
    try:
        return float(s), None
    except ValueError:
        return None, f"변환불가({v})"


def load_aviat213_lookup():
    """Aviat 213종 K코드/KTC코드 매핑 로드 (코드매핑 검증용)."""
    wb = openpyxl.load_workbook(AVIAT213_SRC, data_only=True)
    ws = wb["등록요청"]
    kcodes = set()
    ktc_to_k = {}
    for r in range(4, 217):
        k = ws.cell(row=r, column=17).value
        ktc = ws.cell(row=r, column=18).value
        if k:
            kcodes.add(str(k).strip())
        if ktc and k:
            ktc_to_k[str(ktc).strip()] = str(k).strip()
    return kcodes, ktc_to_k


def main():
    wb_src = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb_src["주문현황"]
    aviat_kcodes, ktc_to_k = load_aviat213_lookup()

    records = []
    date_errors = []       # (row, field, raw, reason)
    numeric_errors = []    # (row, field, raw, reason)
    amount_mismatches = []
    code_map_issues = []
    dup_keys = Counter()

    for r in range(DATA_START, DATA_END + 1):
        def g(name):
            return ws.cell(row=r, column=C[name]).value

        order_no = g("주문번호")
        item_no = g("품목번호")
        combo_key = f"{order_no}_{item_no}"
        dup_keys[combo_key] += 1

        order_date, e1 = parse_date(g("주문일자"))
        if e1:
            date_errors.append((r, "주문일자", g("주문일자"), e1))
        po_date, e2 = parse_date(g("발주일자"))
        if e2:
            date_errors.append((r, "발주일자", g("발주일자"), e2))
        due_date, e3 = parse_date(g("배송희망일"))
        if e3:
            date_errors.append((r, "배송희망일", g("배송희망일"), e3))
        recv_date, e4 = parse_date(g("입고일자"))
        if e4:
            date_errors.append((r, "입고일자", g("입고일자"), e4))

        qty, eq = parse_number(g("수량"))
        if eq:
            numeric_errors.append((r, "수량", g("수량"), eq))
        sale_price, esp = parse_number(g("판매단가"))
        if esp:
            numeric_errors.append((r, "판매단가", g("판매단가"), esp))
        sale_amt, esa = parse_number(g("판매금액"))
        if esa:
            numeric_errors.append((r, "판매금액", g("판매금액"), esa))
        buy_price, ebp = parse_number(g("매입단가"))
        if ebp:
            numeric_errors.append((r, "매입단가", g("매입단가"), ebp))
        buy_amt, eba = parse_number(g("매입금액"))
        if eba:
            numeric_errors.append((r, "매입금액", g("매입금액"), eba))

        if None not in (qty, sale_price, sale_amt) and abs(qty * sale_price - sale_amt) > 1:
            amount_mismatches.append((r, order_no, "판매", qty, sale_price, sale_amt, qty * sale_price))
        if None not in (qty, buy_price, buy_amt) and abs(qty * buy_price - buy_amt) > 1:
            amount_mismatches.append((r, order_no, "매입", qty, buy_price, buy_amt, qty * buy_price))

        kcode_raw = g("K코드")
        kcode_orig = str(kcode_raw).strip() if kcode_raw else ""
        kcode = kcode_orig
        product_code = g("상품코드")
        candidate = ktc_to_k.get(str(product_code).strip()) if product_code else None
        kcode_backfilled = False

        if not kcode and candidate:
            kcode = candidate
            kcode_backfilled = True
            code_map_issues.append((
                r, order_no, product_code, "",
                f"고객사 상품코드(K코드) 공란 → 상품코드({product_code})를 Aviat 213종 KTC코드와 매칭하여 "
                f"K코드 '{kcode}'로 역추정(백필). 원본 파일에는 반영하지 않고 본 정제 결과에만 표시.",
                "추정(KTC코드 1:1 매칭)",
            ))
        elif not kcode:
            code_map_issues.append((
                r, order_no, product_code, "", "고객사 상품코드(K코드) 공란, 상품코드로도 매핑 불가",
                "확인 필요",
            ))
        elif kcode not in aviat_kcodes:
            if candidate and candidate in aviat_kcodes:
                # 1단계에서 확인한 '구코드->신코드' 재부여 4건과 동일 패턴(상품코드는 같고 K코드만 다름)
                code_map_issues.append((
                    r, order_no, product_code, kcode_orig,
                    f"발주이력 K코드('{kcode_orig}')가 현재 계약목록에 없으나, 상품코드({product_code}) 기준 "
                    f"Aviat 213종 KTC코드와 매칭하면 K코드 '{candidate}'로 대응됨. 1단계에서 확인한 "
                    f"'계약 갱신 시 K코드 재부여' 4건(K9090176/K9093723/K9188156/K9210278~281)과 동일 "
                    f"패턴으로, 발주 당시에는 구K코드로 기록되었을 가능성이 높음.",
                    "추정(상품코드 1:1 매칭, 1단계 재부여 4건과 일치)",
                ))
                kcode = candidate
            else:
                code_map_issues.append((
                    r, order_no, product_code, kcode,
                    "발주이력의 K코드가 Aviat 213종 계약품목 목록에 없음(계약외 품목이거나 코드 변경 가능성)",
                    "확인 필요",
                ))

        returned_or_exchanged = bool(g("반품/교환") and str(g("반품/교환")).strip())
        cancel_qty, _ = parse_number(g("취소수량"))
        return_qty, _ = parse_number(g("반품수량"))
        is_cancel_or_return = returned_or_exchanged or (cancel_qty or 0) > 0 or (return_qty or 0) > 0
        recv_qty, _ = parse_number(g("입고수량"))
        fully_received = (g("주문상태") == "전량입고") or (recv_qty is not None and qty is not None and recv_qty >= qty)

        records.append({
            "행": r, "결합키": combo_key, "주문번호": order_no, "품목번호": item_no,
            "주문일자": order_date, "발주일자": po_date, "배송희망일": due_date, "입고일자": recv_date,
            "주문자": g("주문자"), "부서": g("부서"),
            "배송지": (g("배송지") or "").strip(), "배송지상세": (g("배송지상세") or "").strip(),
            "상품코드": product_code, "K코드": kcode, "K코드_원본기록값": kcode_orig,
            "K코드_백필여부": kcode_backfilled, "K코드_보정여부": kcode != kcode_orig and not kcode_backfilled,
            "품명": g("상품명"), "규격": g("대표규격"), "제조사": g("제조사"), "단위": g("단위"),
            "수량": qty, "판매단가": sale_price, "판매금액": sale_amt,
            "매입단가": buy_price, "매입금액": buy_amt,
            "주문상태": g("주문상태"), "반품교환여부": "Y" if returned_or_exchanged else "N",
            "취소수량": cancel_qty, "반품수량": return_qty, "전량입고여부": "Y" if fully_received else "N",
            "매출구분": g("매출구분"), "빌딩명": g("빌딩명"), "국사명": g("국사명"), "프로젝트명": g("프로젝트명"),
            "제외후보": is_cancel_or_return,
        })

    dup_combo = {k: c for k, c in dup_keys.items() if c > 1}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------------- 정제데이터 ----------------
    ws1 = wb.create_sheet("정제데이터")
    headers = ["결합키(주문번호_품목번호)", "주문번호", "품목번호", "주문일자", "발주일자", "배송희망일",
               "입고일자", "주문자", "부서", "배송지", "배송지상세", "상품코드", "K코드(보정반영)",
               "K코드_원본기록값", "K코드_처리방식", "품명", "규격", "제조사", "단위", "수량", "판매단가",
               "판매금액", "매입단가", "매입금액", "주문상태", "반품교환여부", "취소수량", "반품수량",
               "전량입고여부", "매출구분", "빌딩명", "국사명", "프로젝트명"]
    ws1.append(headers)
    for it in records:
        r = ws1.max_row + 1
        if it["K코드_백필여부"]:
            method = "백필(상품코드 매칭)"
        elif it["K코드_보정여부"]:
            method = "보정(구코드→신코드)"
        else:
            method = "원본"
        ws1.append([
            it["결합키"], it["주문번호"], it["품목번호"], it["주문일자"], it["발주일자"], it["배송희망일"],
            it["입고일자"], it["주문자"], it["부서"], it["배송지"], it["배송지상세"], it["상품코드"],
            it["K코드"], it["K코드_원본기록값"], method, it["품명"], it["규격"], it["제조사"],
            it["단위"], it["수량"], it["판매단가"], it["판매금액"], it["매입단가"], it["매입금액"],
            it["주문상태"], it["반품교환여부"], it["취소수량"], it["반품수량"], it["전량입고여부"],
            it["매출구분"], it["빌딩명"], it["국사명"], it["프로젝트명"],
        ])
        for c in [2, 3, 8, 9, 10, 11, 12, 14, 16, 17, 18, 19, 20, 21, 23]:
            mark_font(ws1, r, c, ORIGIN_FONT)
        for c in [4, 5, 6, 7]:
            ws1.cell(row=r, column=c).number_format = FMT_DATE
        if it["K코드_백필여부"] or it["K코드_보정여부"]:
            mark_fill(ws1, r, 13, CONFIRMED_FILL)
            mark_fill(ws1, r, 15, CONFIRMED_FILL)
        if it["제외후보"]:
            mark_fill(ws1, r, 26, REVIEW_FILL)
            mark_fill(ws1, r, 28, REVIEW_FILL)
    for c in [21, 22, 23, 24]:
        for r in range(2, ws1.max_row + 1):
            ws1.cell(row=r, column=c).number_format = FMT_AMOUNT
    finalize_sheet(ws1, 1, len(headers), ws1.max_row)

    # ---------------- 제외검토 ----------------
    ws2 = wb.create_sheet("제외검토")
    headers2 = ["결합키", "주문번호", "주문일자", "품명", "K코드", "수량", "사유", "권고"]
    ws2.append(headers2)
    excl_count = 0
    for it in records:
        if it["제외후보"]:
            excl_count += 1
            ws2.append([
                it["결합키"], it["주문번호"], it["주문일자"], it["품명"], it["K코드"], it["수량"],
                f"반품/교환여부={it['반품교환여부']}, 취소수량={it['취소수량']}, 반품수량={it['반품수량']}",
                "구성(BoM) 추정 시 이 행 제외 검토",
            ])
    if excl_count == 0:
        ws2.append(["-", "-", "-", "-", "-", "-", "취소/반품/교환 건 없음(245건 전량 정상)", "제외 대상 없음"])
    for c in [3]:
        for r in range(2, ws2.max_row + 1):
            ws2.cell(row=r, column=c).number_format = FMT_DATE
    finalize_sheet(ws2, 1, len(headers2), ws2.max_row)

    # ---------------- 금액오류 ----------------
    ws3 = wb.create_sheet("금액오류")
    headers3 = ["행", "주문번호", "구분", "수량", "단가", "기록된금액", "수량×단가", "차이"]
    ws3.append(headers3)
    for r, order_no, kind, qty, price, amt, calc in amount_mismatches:
        ws3.append([r, order_no, kind, qty, price, amt, calc, amt - calc])
        mark_fill(ws3, ws3.max_row, 8, ERROR_FILL)
    date_num_error_start_row = ws3.max_row + 2
    if date_errors or numeric_errors:
        ws3.append([])
        ws3.append(["※ 날짜/숫자 변환 오류 (삭제하지 않고 아래에 기록)"])
        ws3.append(["행", "필드", "원본값", "사유"])
        for r, field, raw, reason in date_errors + numeric_errors:
            ws3.append([r, field, raw, reason])
    if not amount_mismatches and not date_errors and not numeric_errors:
        ws3.append(["-", "-", "-", "-", "-", "-", "-", "-"])
        ws3.append(["금액 불일치 0건, 날짜/숫자 변환 오류 0건 (245건 전량 정상)"])
    finalize_sheet(ws3, 1, len(headers3), max(ws3.max_row, 2))

    # ---------------- 코드매핑오류 ----------------
    ws4 = wb.create_sheet("코드매핑오류")
    headers4 = ["행", "주문번호", "상품코드", "K코드(원본)", "이슈", "확정수준"]
    ws4.append(headers4)
    for r, order_no, product_code, kcode, issue, level in code_map_issues:
        ws4.append([r, order_no, product_code, kcode, issue, level])
        mark_fill(ws4, ws4.max_row, 5, ERROR_FILL if level == "확인 필요" else REVIEW_FILL)
    n_all = len(records)
    n_mapped_raw = sum(1 for it in records if it["K코드_원본기록값"] and it["K코드_원본기록값"] in aviat_kcodes)
    n_mapped_final = sum(1 for it in records if it["K코드"] and it["K코드"] in aviat_kcodes)
    n_corrected = sum(1 for it in records if it["K코드_보정여부"])
    n_backfilled = sum(1 for it in records if it["K코드_백필여부"])
    mapping_rate_raw = n_mapped_raw / n_all if n_all else 0
    mapping_rate_final = n_mapped_final / n_all if n_all else 0
    ws4.append([])
    ws4.append([f"K코드 매핑률(원본 그대로, 보정 전, 확정): {n_mapped_raw}/{n_all} = {mapping_rate_raw:.1%}"])
    ws4.append([
        f"K코드 매핑률(백필 {n_backfilled}건 + 구코드→신코드 보정 {n_corrected}건 반영, 추정 포함): "
        f"{n_mapped_final}/{n_all} = {mapping_rate_final:.1%}"
    ])
    ws4.append([
        "※ 보정 {}건은 상품코드(KTC코드) 1:1 매칭에 근거한 '추정'이며, 발주 시점에 실제로 구K코드가 "
        "사용되었음을 공급사·KT 양측 문서로 확정한 것은 아님(확인 필요).".format(n_corrected + n_backfilled)
    ])
    if not code_map_issues:
        ws4.append(["코드 매핑 이슈 없음"])
    finalize_sheet(ws4, 1, len(headers4), max(ws4.max_row, 2))

    # ---------------- 주문요약 ----------------
    ws5 = wb.create_sheet("주문요약")
    ws5.append(["항목", "값"])
    order_dates = [it["주문일자"] for it in records if it["주문일자"]]
    total_sale = sum(it["판매금액"] or 0 for it in records)
    total_buy = sum(it["매입금액"] or 0 for it in records)
    orderer_dist = Counter(it["주문자"] for it in records)
    dept_dist = Counter(it["부서"] for it in records)
    addr_dist = Counter(it["배송지"] for it in records)
    revtype_dist = Counter(it["매출구분"] for it in records)
    summary_rows = [
        ("총 처리 행 수", len(records)),
        ("고유 주문번호 수", len(set(it["주문번호"] for it in records))),
        ("결합키(주문번호_품목번호) 중복 건수", len(dup_combo)),
        ("주문일자 범위", f"{min(order_dates)} ~ {max(order_dates)}" if order_dates else "N/A"),
        ("판매금액 합계", total_sale),
        ("매입금액 합계", total_buy),
        ("취소/반품/교환 제외후보 건수", excl_count),
        ("K코드 매핑률(원본, 보정 전, 확정)", f"{mapping_rate_raw:.1%}"),
        ("K코드 매핑률(백필+보정 반영, 추정 포함)", f"{mapping_rate_final:.1%}"),
        ("코드매핑 확인필요 건수", sum(1 for _, _, _, _, _, lv in code_map_issues if lv == "확인 필요")),
        ("K코드 백필(추정) 건수", n_backfilled),
        ("K코드 구코드→신코드 보정(추정) 건수", n_corrected),
        ("주문자 고유값 수", len(orderer_dist)),
        ("주문자 분포", ", ".join(f"{k}:{v}건" for k, v in orderer_dist.most_common())),
        ("부서 고유값 수", len(dept_dist)),
        ("부서 분포", ", ".join(f"{k}:{v}건" for k, v in dept_dist.most_common())),
        ("배송지 고유값 수", len(addr_dist)),
        ("배송지 분포(상위)", " | ".join(f"{k[:30]}...:{v}건" for k, v in addr_dist.most_common(5))),
        ("매출구분 분포", ", ".join(f"{k}:{v}건" for k, v in revtype_dist.most_common())),
        ("빌딩명/국사명/프로젝트명 채움 건수", sum(1 for it in records if it["빌딩명"] or it["국사명"] or it["프로젝트명"])),
    ]
    for k, v in summary_rows:
        ws5.append([k, v])
    ws5.append([])
    ws5.append(["※ 주의사항(확인 필요)", ""])
    ws5.append([
        "주문자 표기 불일치 의심",
        "'이상*'(188건)과 '이상헌'(42건)은 이메일 패턴('sangheon')상 동일인일 가능성이 있으나 "
        "부서가 각각 '그룹시설팀'/'강남유선시설2팀'으로 다르게 기록되어 있어 임의로 병합하지 않음. "
        "동일인이 부서 이동했을 가능성과, 단순 우연의 일치일 가능성을 모두 열어둠(확인 필요).",
    ])
    ws5.append([
        "빌딩명/국사명/프로젝트명 전부 공란",
        "245건 전체에서 현장 식별 보조 필드(빌딩명/국사명/프로젝트명)가 100% 공란. "
        "실제 설치 국소를 배송지만으로 특정할 수 없어, 4단계 이벤트 분석에서 '배송지'는 "
        "참고 신호로만 사용하고 단독으로 세트를 확정하지 않음.",
    ])
    ws5.append([
        "대량 동일 배송지 반복",
        "'경기 안양시 ***'(마스킹) 배송지가 188건, 11개월에 걸쳐 반복 등장 → 특정 링크 설치가 아니라 "
        "본사/창고성 집하지로 추정(확인 필요). 동일 배송지라는 이유만으로 하나의 세트로 묶지 않음.",
    ])
    ws5.append([
        "주문번호 체계 상이 2건",
        "'MR'로 시작하는 주문번호 2건(MR20250901-71950, MR20250704-44811)은 나머지 243건('45'로 "
        "시작하는 SAP식 번호)과 체계가 달라 K코드 필드가 공란이었음 → 상품코드 기준으로 K코드 백필.",
    ])
    finalize_sheet(ws5, 1, 2, ws5.max_row)
    ws5.column_dimensions["B"].width = 90

    out_path = OUTPUT_DIR / "02_Aviat_발주이력_정제.xlsx"
    wb.save(out_path)

    print("=== 3단계: Aviat 발주이력 정제 완료 ===")
    print(f"입력 파일: {SRC.name} (시트 '주문현황' row{DATA_START}~{DATA_END})")
    print(f"처리 행 수: {len(records)}건")
    print(f"제외 또는 오류 행 수: 0건 삭제(취소/반품 후보 {excl_count}건은 표시만 하고 유지)")
    print(f"생성 파일: {out_path}")
    print("주요 검증 결과:")
    print(f"  - 결합키 중복: {len(dup_combo)}건")
    print(f"  - 날짜 변환 오류: {len(date_errors)}건, 숫자 변환 오류: {len(numeric_errors)}건")
    print(f"  - 금액(단가×수량≠기록금액) 불일치: {len(amount_mismatches)}건")
    print(f"  - K코드 매핑률(원본, 확정): {n_mapped_raw}/{n_all} = {mapping_rate_raw:.1%}")
    print(f"  - K코드 매핑률(백필{n_backfilled}건+보정{n_corrected}건 반영, 추정 포함): {n_mapped_final}/{n_all} = {mapping_rate_final:.1%}")
    print(f"  - 코드매핑 확인필요: {sum(1 for _,_,_,_,_,lv in code_map_issues if lv=='확인 필요')}건")
    print(f"  - 취소/반품/교환 후보: {excl_count}건")


if __name__ == "__main__":
    main()
