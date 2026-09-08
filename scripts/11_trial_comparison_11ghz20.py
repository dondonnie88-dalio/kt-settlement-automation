"""
회신 전 시범 비교(돈현님 요청, 2026-09-04, 4차): 11GHz 2+0만 먼저 Ceragon과 나란히 놓아
"비교가 실제로 어떻게 보이는지" 감을 잡는다.
입력: output/04_Aviat_구성방식별_표준BoM_추정_v2.xlsx, output/06_Ceragon_구성방식별_BoM_정리.xlsx
출력: output/11_11GHz_2+0_시범비교(회신전_참고용).xlsx

주의: 이 파일은 K코드 매칭이 아니라 '규모 감각'만 보기 위한 참고 자료다. Aviat 쪽은 04v2에서
독립 반복으로 확인된 유력 후보 6개 품목만 담고 있어(케이블·인터페이스카드·전원장치 등은 아직
반복 근거가 없어 제외) 완성된 링크 BoM이 아니며, Ceragon 쪽은 케이블·커넥터까지 포함한 완성
링크 총액이다. 두 총액을 직접 뺄셈하거나 협상 근거로 쓰지 않는다 - 애니콤 회신을 받아 07번
파일의 정식 비교로 대체해야 한다.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_style import finalize_sheet, REVIEW_FILL, CONFIRMED_FILL, FMT_AMOUNT, mark_fill

BASE_DIR = Path("/home/user/kt-settlement-automation")
OUTPUT_DIR = BASE_DIR / "output"
SRC_BOM_V2 = OUTPUT_DIR / "04_Aviat_구성방식별_표준BoM_추정_v2.xlsx"
SRC_CERAGON = OUTPUT_DIR / "06_Ceragon_구성방식별_BoM_정리.xlsx"


def load_aviat_strong_items():
    wb = openpyxl.load_workbook(SRC_BOM_V2, data_only=True)
    ws = wb["11GHz_IAP3_2+0"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["판정수준"]] != "유력 후보":
            continue
        # 2026-09-08 정정: 이 열은 N당 배수(비율)가 아니라 링크 전체(양쪽 사이트 합산) 실제
        # 수량이다 - 이전 버전은 배수를 그대로 써서 2+0 기준 실제 수량의 절반으로 축소되어 있었음.
        qty = row[idx["추정 수량(N+0 링크 전체 가정, 양쪽 사이트 합산)"]]
        price = row[idx["판매단가"]] or 0
        rows.append({
            "K코드": row[idx["K코드"]], "품명": row[idx["품명"]], "품목군": row[idx["품목군"]],
            "수량": qty, "판매단가": price, "판매금액": qty * price if isinstance(qty, (int, float)) else None,
        })
    return rows


def load_ceragon_items(sd):
    wb = openpyxl.load_workbook(SRC_CERAGON, data_only=True)
    ws = wb["BoM_상세"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["원본주파수"]] != "11GHz" or row[idx["표준화구성표기"]] != "2+0" or row[idx["SD구분"]] != sd:
            continue
        rows.append({
            "K코드": row[idx["K코드"]], "품명": row[idx["품명"]],
            "수량": row[idx["1개링크_총수량"]], "인하단가": row[idx["최종인하단가"]],
            "인하금액": row[idx["인하금액"]],
        })
    return rows


def main():
    aviat_rows = load_aviat_strong_items()
    ceragon_nonsd = load_ceragon_items("Non_SD")
    ceragon_sd = load_ceragon_items("SD")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ================= 안내 =================
    ws0 = wb.create_sheet("안내")
    ws0.append(["항목", "내용"])
    guide = [
        ("목적", "애니콤 회신이 오기 전, 11GHz 2+0 하나만 먼저 Ceragon과 나란히 놓아 비교가 실제로 "
                "어떻게 보이는지 감을 잡기 위한 시범 자료다(돈현님 요청, 2026-09-04)."),
        ("이 파일로 하지 말아야 할 것", "이 파일의 두 총액을 빼거나 비율을 내서 '이만큼 비싸다/싸다'는 "
                                  "결론을 내리지 않는다. 협상 자료나 최종 보고서에 인용하지 않는다."),
        ("왜 직접 비교가 안 되는가 (1) 품목 범위", "Aviat 쪽은 04v2에서 서로 다른 2개의 독립 이벤트로 "
                                            "반복 확인된 '유력 후보' 6개 품목만 담았다(VR4 CHASSIS/"
                                            "ODU LOW·HIGH/FAN-CV/Mounting Bracket/Node License). "
                                            "실제 링크에는 모뎀·인터페이스카드·케이블·전원장치·SFP도 "
                                            "필요하지만 이런 품목은 아직 반복 근거가 없어 '참고 후보'로 "
                                            "남아 있어 제외했다 - 즉 Aviat 쪽 수치는 처음부터 '완성된 "
                                            "링크'가 아니라 '핵심 뼈대'만이다."),
        ("왜 직접 비교가 안 되는가 (2) 구성 방식 차이", "07번 파일 품목구성_차이 시트 기준: Aviat는 "
                                              "모듈(카드) 추가형(IDU 셀프에 카드 추가), Ceragon은 "
                                              "시스템(셀프) 단위 - 같은 '2+0'이라도 부품 구성 자체가 "
                                              "다르다."),
        ("왜 직접 비교가 안 되는가 (3) K코드 매칭 불가", "SFP류 3건(직접대응_품목 시트)을 제외하면 "
                                               "두 공급사의 K코드는 서로 대응되지 않는다. 이 표는 "
                                               "품목을 코드로 맞춘 것이 아니라 각자 목록을 그대로 "
                                               "나열한 것이다."),
        ("정식 비교로 가는 길", "애니콤 회신이 오면 (1) 10번 대조표로 Aviat 자체 회신 대 발주이력 "
                          "추정을 먼저 맞추고 (2) 그렇게 확정된 Aviat 표준 BoM을 07번 파일의 "
                          "구성총액_비교 시트에 반영해 Ceragon과 정식 비교한다. 이 파일은 그 전 "
                          "단계의 참고용일 뿐이다."),
    ]
    for k, v in guide:
        ws0.append([k, v])
    finalize_sheet(ws0, 1, 2, ws0.max_row)
    ws0.column_dimensions["B"].width = 100

    # ================= Aviat 11GHz 2+0(유력후보) =================
    ws1 = wb.create_sheet("Aviat_11GHz_2+0(유력후보)")
    headers1 = ["K코드", "품명", "품목군", "수량(1개 링크)", "판매단가", "판매금액"]
    ws1.append(headers1)
    total_aviat = 0
    for row in aviat_rows:
        r = ws1.max_row + 1
        ws1.append([row["K코드"], row["품명"], row["품목군"], row["수량"], row["판매단가"], row["판매금액"]])
        mark_fill(ws1, r, 1, CONFIRMED_FILL)
        if row["판매금액"]:
            total_aviat += row["판매금액"]
    r = ws1.max_row + 1
    ws1.append(["", "소계(유력 후보 6개 품목만, 완성 링크 아님)", "", "", "", total_aviat])
    for c in [5, 6]:
        for rr in range(2, ws1.max_row + 1):
            ws1.cell(row=rr, column=c).number_format = FMT_AMOUNT
    finalize_sheet(ws1, 1, len(headers1), ws1.max_row)

    # ================= Ceragon 11GHz 2+0 (Non_SD / SD) =================
    ceragon_totals = {}
    for sd, rows in [("Non_SD", ceragon_nonsd), ("SD", ceragon_sd)]:
        ws = wb.create_sheet(f"Ceragon_11GHz_2+0_{sd}")
        headers = ["K코드", "품명", "수량(1개 링크 총수량)", "인하단가", "인하금액"]
        ws.append(headers)
        total = 0
        for row in rows:
            ws.append([row["K코드"], row["품명"], row["수량"], row["인하단가"], row["인하금액"]])
            total += row["인하금액"] or 0
        ws.append(["", f"소계({len(rows)}개 품목, 완성 링크 BoM)", "", "", total])
        for c in [4, 5]:
            for rr in range(2, ws.max_row + 1):
                ws.cell(row=rr, column=c).number_format = FMT_AMOUNT
        finalize_sheet(ws, 1, len(headers), ws.max_row)
        ceragon_totals[sd] = (len(rows), total)

    # ================= 요약비교 =================
    ws_sum = wb.create_sheet("요약비교(참고용)")
    ws_sum.append(["구분", "품목수", "총액(참고용)", "포함 범위"])
    ws_sum.append(["Aviat 11GHz 2+0", len(aviat_rows), total_aviat,
                    "유력 후보 6개 품목만(모뎀·인터페이스카드·케이블·전원 등 미포함) - 완성 링크 아님"])
    for sd in ["Non_SD", "SD"]:
        n, total = ceragon_totals[sd]
        ws_sum.append([f"Ceragon 11GHz 2+0 ({sd})", n, total, "케이블·커넥터·전원까지 포함한 완성 링크 BoM"])
    r = ws_sum.max_row + 1
    ws_sum.append(["주의", "", "", "위 총액은 포함 범위가 서로 달라 직접 빼거나 나눠서 '몇 배 "
                             "차이'라고 말할 수 없다. 애니콤 회신으로 Aviat 쪽을 완성 BoM으로 "
                             "채운 뒤 07번 파일에서 다시 비교해야 한다."])
    mark_fill(ws_sum, r, 1, REVIEW_FILL)
    mark_fill(ws_sum, r, 4, REVIEW_FILL)
    for rr in range(2, ws_sum.max_row):
        ws_sum.cell(row=rr, column=3).number_format = FMT_AMOUNT
    finalize_sheet(ws_sum, 1, 4, ws_sum.max_row)
    ws_sum.column_dimensions["D"].width = 60

    sheet_order = ["안내", "Aviat_11GHz_2+0(유력후보)", "Ceragon_11GHz_2+0_Non_SD",
                   "Ceragon_11GHz_2+0_SD", "요약비교(참고용)"]
    wb._sheets = [wb[name] for name in sheet_order]

    out_path = OUTPUT_DIR / "11_11GHz_2+0_시범비교(회신전_참고용).xlsx"
    wb.save(out_path)

    print("=== 회신 전 시범 비교(11GHz 2+0) 생성 완료 ===")
    print(f"Aviat 11GHz 2+0 유력 후보: {len(aviat_rows)}개 품목, 참고 판매총액 {total_aviat:,}원")
    for sd in ["Non_SD", "SD"]:
        n, total = ceragon_totals[sd]
        print(f"Ceragon 11GHz 2+0 ({sd}): {n}개 품목, 인하금액합계 {total:,.0f}원")
    print(f"생성 파일: {out_path}")
    print("※ 참고용 시범 비교입니다. 애니콤 정식 회신 전까지 협상/보고서 근거로 사용하지 마세요.")


if __name__ == "__main__":
    main()
