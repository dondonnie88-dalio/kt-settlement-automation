"""
18단계: 대성인포텍向 신규 SW패키지 SD/Non_SD 스케일링 확인 요청 (돈현님 요청, 2026-09-15)

배경: "신규품목 하나도 안 깎고 합의했는데 실수한 거 아닐까요?"라는 질문에서 출발.
19개 신규 SW패키지 품목(K9212896, K9213016~033)은 '기존단가(정가)' 자체가 없어
16번의 마진역산 로직(정가 x (1-마진율))을 적용할 근거가 없었다 - 42건 협상 대상에서
빠진 건 누락이 아니라 애초에 비교할 기준값이 없었기 때문.

다만 이번 8/11GHz 딜에 포함된 8개 품목(2+0/4+0/6+0/8+0 x SD/Non_SD)의 가격을
살펴보면, 6+0/8+0 구간에서 SD가 Non_SD보다 오히려 싸다(2+0/4+0은 정상적으로 SD가
더 비쌈). 카탈로그 내 실제 SD/Non_SD 짝(ICC/ICB 8쌍, 전부 정가 있음)은 예외 없이
SD가 31~39% 비싸다는 게 유일하게 확인 가능한 근거인데, 이 패턴과 반대 방향이다.

이건 "얼마가 맞다"고 KT가 판단할 근거가 없는 사안(정가도 없고 신규 SKU라 비교할 과거
이력도 없음) - 대성에 "의도된 스케일링이 맞는지" 확인을 요청하는 것으로, 이미 합의된
42건 협상과는 무관한 별개 건이다. 확인 전까지 17번 파일에서 해당 4개(8/11GHz 6+0/8+0
SD·Non_SD)는 '확인 중(잠정치)'로 표시해 뒀다.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_style import finalize_sheet, ERROR_FILL, REVIEW_FILL, FMT_AMOUNT, WRAP_TOP, mark_fill

BASE_DIR = Path("/home/user/kt-settlement-automation")
OUTPUT_DIR = BASE_DIR / "output"
SRC_CERAGON = OUTPUT_DIR / "06_Ceragon_구성방식별_BoM_정리.xlsx"

# 카탈로그 내 실제 SD/Non_SD 짝(둘 다 정가 있음) - 품명 매칭용
REFERENCE_PAIRS = [
    ("CER_32T-ICC-6L", "CER_32T-ICC-SD-6L"),
    ("CER_32T-ICC-6H", "CER_32T-ICC-SD-6H"),
    ("CER_32T-ICC-08", "CER_32T-ICC-SD-08"),
    ("CER_32T-ICC-11", "CER_32T-ICC-SD-11"),
    ("CER_32T-ICB-6L", "CER_32T-ICB-SD-6L"),
    ("CER_32T-ICB-6H", "CER_32T-ICB-SD-6H"),
    ("CER_32T-ICB-08", "CER_32T-ICB-SD-08"),
    ("CER_32T-ICB-11", "CER_32T-ICB-SD-11"),
]

# 이번 딜(8/11GHz) 관련 4개 - 확인 요청 핵심 대상
DEAL_PAIRS = [
    ("2+0", "K9213024", "K9213023"),
    ("4+0", "K9213026", "K9213025"),
    ("6+0", "K9213028", "K9213027"),
    ("8+0", "K9213030", "K9213029"),
]
# 참고용(6GHz, 이번 딜과 무관하지만 동일 패턴 확인용)
REF_BAND_PAIRS = [
    ("2+0", "K9213016", "K9212896"),
    ("4+0", "K9213018", "K9213017"),
    ("6+0", "K9213020", "K9213019"),
    ("8+0", "K9213022", "K9213021"),
]


def load_master():
    wb = openpyxl.load_workbook(SRC_CERAGON, data_only=True)
    ws = wb["품목마스터"]
    h = [c.value for c in ws[1]]
    idx = {v: k for k, v in enumerate(h)}
    by_code = {r[idx["K코드"]]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    by_name = {r[idx["품명"]].strip(): r for r in ws.iter_rows(min_row=2, values_only=True)}
    return by_code, by_name, idx


def build_reference_rows():
    _, by_name, idx = load_master()
    rows = []
    for plain, sd in REFERENCE_PAIRS:
        p, s = by_name[plain], by_name[sd]
        pp, sp = p[idx["최종인하단가"]], s[idx["최종인하단가"]]
        rows.append((plain, pp, sd, sp, sp / pp))
    return rows


def build_deal_rows(pairs, band_label):
    by_code, _, idx = load_master()
    rows = []
    for cfg, nonsd_code, sd_code in pairs:
        n, s = by_code[nonsd_code], by_code[sd_code]
        np_, sp = n[idx["최종인하단가"]], s[idx["최종인하단가"]]
        rows.append((band_label, cfg, nonsd_code, n[idx["품명"]].strip(), np_,
                     sd_code, s[idx["품명"]].strip(), sp, sp / np_))
    return rows


def build_message_sheet(wb, ref_rows, deal_rows):
    ws = wb.create_sheet("확인요청_초안")
    ws.append(["항목", "내용"])

    ratio_lo = min(r[4] for r in ref_rows)
    ratio_hi = max(r[4] for r in ref_rows)

    lines = "\n".join(
        f"  - {cfg}: Non_SD {npv:,.0f}원 vs SD {spv:,.0f}원 (SD/Non_SD = {ratio:.2f})"
        for _, cfg, _, _, npv, _, _, spv, ratio in deal_rows
    )

    msg = (
        "안녕하세요, KT commerce 이돈현입니다.\n\n"
        "귀사에서 신규로 제공해 주신 8GHz/11GHz Configuration SW Package 단가 중 "
        "확인 부탁드릴 부분이 있어 문의드립니다.\n\n"
        f"귀사 단가표 내 실제 SD/Non_SD 짝(ICC/ICB 계열 8쌍)은 예외 없이 SD가 "
        f"Non_SD보다 {ratio_lo-1:.0%}~{ratio_hi-1:.0%} 더 높게 책정되어 있습니다. "
        "그런데 이번에 제공해 주신 8GHz/11GHz Configuration SW Package 중 "
        "6+0/8+0 구성에서는 아래와 같이 SD가 오히려 Non_SD보다 낮게 책정되어 "
        "있습니다(2+0/4+0 구성은 다른 품목들과 동일하게 SD가 더 높음).\n\n"
        + lines + "\n\n"
        "의도된 가격 정책인지, 혹은 수정이 필요한 부분인지 확인 부탁드립니다. "
        "참고로 6GHz Configuration SW Package에서도 동일한 패턴이 확인되어 "
        "함께 첨부(비교근거_상세 시트)해 드립니다.\n\n"
        "감사합니다.\n"
        "KT commerce 이돈현 드림"
    )

    guide = [
        ("용도", "대성인포텍에 보낼 이메일 본문 초안. '비교근거_상세' 시트를 첨부하거나 "
                "표를 붙여넣으면 된다."),
        ("메시지 초안", msg),
        ("성격(내부용)", "이미 합의된 16번의 42건 협상과는 무관한 별개 건 - 이 4개 품목은 "
                "정가 자체가 없어 42건 협상 대상에 포함될 수 없었다. '얼마가 맞다'고 KT가 "
                "주장하는 게 아니라, 카탈로그 내 유일하게 확인 가능한 SD 프리미엄 근거"
                "(8쌍, 31~39%)와 반대 방향인 걸 짚어 확인을 요청하는 것 - 근거가 하드웨어"
                "(ICC/ICB)와 SW 라이선스 간 비교라는 한계는 있음."),
        ("주의", "이 초안은 자동 생성된 것이니 보내시기 전에 실제 상황에 맞게 다듬어서 "
              "사용하시기 바랍니다."),
    ]
    for k, v in guide:
        ws.append([k, v])
    finalize_sheet(ws, 1, 2, ws.max_row)
    ws.column_dimensions["B"].width = 100
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).alignment = WRAP_TOP


def build_detail_sheet(wb, ref_rows, deal_rows, ref_band_rows):
    ws = wb.create_sheet("비교근거_상세")
    ws.append(["구분", "Non_SD 품명/K코드", "Non_SD 단가", "SD 품명/K코드", "SD 단가", "SD/Non_SD 배율"])

    ws.append(["=== 카탈로그 내 실제 SD/Non_SD 짝(정가 있음, 참고 기준) ===", None, None, None, None, None])
    mark_fill(ws, ws.max_row, 1, REVIEW_FILL)
    for plain, pp, sd, sp, ratio in ref_rows:
        ws.append(["기준(ICC/ICB)", plain, pp, sd, sp, ratio])

    ws.append(["=== 이번 딜(8/11GHz) 신규 SW패키지 - 확인 요청 대상 ===", None, None, None, None, None])
    mark_fill(ws, ws.max_row, 1, ERROR_FILL)
    for band, cfg, ncode, nname, npv, scode, sname, spv, ratio in deal_rows:
        rr = ws.max_row + 1
        ws.append([f"{band} {cfg}", f"{nname}({ncode})", npv, f"{sname}({scode})", spv, ratio])
        if ratio < 1:
            mark_fill(ws, rr, 6, ERROR_FILL)

    ws.append(["=== 6GHz 신규 SW패키지 - 참고(동일 패턴 확인용, 이번 딜과 무관) ===", None, None, None, None, None])
    mark_fill(ws, ws.max_row, 1, REVIEW_FILL)
    for band, cfg, ncode, nname, npv, scode, sname, spv, ratio in ref_band_rows:
        rr = ws.max_row + 1
        ws.append([f"{band} {cfg}", f"{nname}({ncode})", npv, f"{sname}({scode})", spv, ratio])
        if ratio < 1:
            mark_fill(ws, rr, 6, REVIEW_FILL)

    for c in [3, 5]:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = FMT_AMOUNT
    finalize_sheet(ws, 1, 6, ws.max_row)
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["D"].width = 45


def main():
    ref_rows = build_reference_rows()
    deal_rows = build_deal_rows(DEAL_PAIRS, "8/11GHz")
    ref_band_rows = build_deal_rows(REF_BAND_PAIRS, "6GHz")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_message_sheet(wb, ref_rows, deal_rows)
    build_detail_sheet(wb, ref_rows, deal_rows, ref_band_rows)
    wb._sheets = [wb["확인요청_초안"], wb["비교근거_상세"]]

    out_path = OUTPUT_DIR / "18_대성_SD스케일링_확인요청.xlsx"
    wb.save(out_path)

    print("=== 18단계: 대성 SD/Non_SD 스케일링 확인 요청 생성 완료 ===")
    print(f"기준 참고쌍(정가 있음): {len(ref_rows)}건, 배율 범위 "
          f"{min(r[4] for r in ref_rows):.2f}~{max(r[4] for r in ref_rows):.2f}")
    print("확인 요청 대상(8/11GHz, 이번 딜 포함):")
    for _, cfg, ncode, _, npv, scode, _, spv, ratio in deal_rows:
        print(f"  {cfg}: {ncode}={npv:,.0f}원 vs {scode}={spv:,.0f}원 (배율 {ratio:.2f})")
    print(f"생성 파일: {out_path}")


if __name__ == "__main__":
    main()
