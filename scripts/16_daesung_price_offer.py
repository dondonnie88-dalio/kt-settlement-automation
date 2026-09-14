"""
16단계: 대성인포텍(Ceragon) 매입단가 제시 (돈현님 요청, 2026-09-14, 11차)

핵심 전략 전환(2026-09-14, 돈현님 지적): "14%를 할인해주겠냐? 잘 생각해봐" - 이전
버전(10차)은 SFP 3종에 대해 "동일 계열 다른 품목들이 쓰는 14.69% 표준 할인율을
적용해달라"고 요청했는데, 이는 대성 입장에서 할인율을 1%->14.69%로 13.69%p나
올려달라는, 근거는 있지만 폭이 지나치게 큰 요청이었다.

돈현님이 원하는 방식: "문제가 된 품목들의 판매단가를 현재단가(정가)로 하고, 거기서
적정 마진을 남길 수 있게 매입가를 조정" - 즉 대성 자체 가격 패턴이 어떻든 상관없이,
"우리는 이 품목을 정가 그대로 고객에게 팔 것이고, 그 안에서 KT의 통상 마진(카테고리별
2%/3%/5%, 01번 Aviat 데이터로 검증된 값)을 남기려면 이 가격에 매입해야 한다"는
KT 자체 마진구조 논리로 필요 매입가를 역산한다. 대성의 할인 패턴이 정상이냐 이상하냐를
따질 필요가 없어 완전히 다른 근거이며, 대부분 품목(안테나 36종)은 2%->3%로 단 1%p만
올려달라는 요청이 되어 훨씬 현실적이다.

대상 품목: 15번 파일의 '전체카탈로그_단가제시안(493종)' 시트에서 '정가상한 적용'으로
표시된 42개 품목(할인율<KT 마진율이라 판매가가 정가를 초과해버리는 품목) 전부를
가져와 필요매입가 = 정가 x (1-마진율)로 재계산한다. 42개 중 2개(Am-1/2-18-CIRC-CR)는
부동소수점 경계상 이미 필요조건을 충족해(차이 0) 실제 요청에서는 제외.

별도로, 표기 오류로 판단되는 2건(RFUC-CPLR-8, RFUC-TWIST Kit-6 - 할인 후 단가가
정가보다 높고 둘 다 634,418원으로 동일값)은 이 마진역산 로직과 무관하게, 정가가
동일한 형제 품목의 실제 단가를 그대로 근거로 사용(마진역산 최소요구치보다도 낮아
KT에 더 유리하고 형제 품목 실측값이라 근거도 더 강함).
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_style import (
    finalize_sheet, ERROR_FILL, REVIEW_FILL, CONFIRMED_FILL, FMT_AMOUNT, FMT_PERCENT,
    mark_fill, WRAP_TOP,
)

BASE_DIR = Path("/home/user/kt-settlement-automation")
OUTPUT_DIR = BASE_DIR / "output"
SRC_CERAGON = OUTPUT_DIR / "06_Ceragon_구성방식별_BoM_정리.xlsx"
SRC_15 = OUTPUT_DIR / "15_매입판매가_제시안.xlsx"

GAP_THRESHOLD = 1  # 원 - 부동소수점 경계로 인한 가짜 차이 제거

# 표기오류 2건의 대체 단가 근거: 동일 정가를 가진 형제 품목의 실제 적용 단가
SIBLING_REF = {
    "K9178851": ("K9178850", "CER_RFUC-CPLR-6"),       # 정가 548,550원 동일
    "K9178856": ("K9178857", "CER_RFUC-TWIST Kit-8"),  # 정가 123,210원 동일
}


def load_master():
    wb = openpyxl.load_workbook(SRC_CERAGON, data_only=True)
    ws = wb["품목마스터"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_code = {row[idx["K코드"]]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    return by_code, idx


def load_margin_driven_items():
    """15번의 '정가상한 적용' 42건 - 판매가를 정가로 고정할 때 KT 마진율을 남기려면
    필요한 매입가(정가 x (1-마진율))를 계산한다."""
    wb15 = openpyxl.load_workbook(SRC_15, data_only=True)
    ws = wb15["전체카탈로그_단가제시안(493종)"]
    h = [c.value for c in ws[1]]
    i = {v: k for k, v in enumerate(h)}
    by_code, idx06 = load_master()

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        flag = r[i["가격 이상 여부"]]
        if not flag or "정가상한" not in str(flag):
            continue
        kcode = r[i["K코드"]]
        margin = r[i["적용 마진율"]]
        cur_price = r[i["매입단가(대성 인하단가)"]]
        list_price = by_code[kcode][idx06["기존단가"]]
        need_price = list_price * (1 - margin)
        gap = cur_price - need_price
        if gap <= GAP_THRESHOLD:
            continue  # 이미 마진 조건 충족(부동소수점 경계) - 요청 불필요
        rows.append({
            "K코드": kcode, "품명": r[i["품명"]].strip(), "마진분류": r[i["마진 분류"]],
            "마진율": margin, "정가": list_price, "대성_현재제시단가": cur_price,
            "대성_현재할인율": 1 - cur_price / list_price,
            "KT_제시매입단가": round(need_price), "차액": gap,
            "구분": "마진확보", "근거": (f"판매단가를 귀사 정가로 유지하며 KT 통상 마진율 "
                                    f"{margin:.0%}를 확보하기 위한 필요 매입단가"),
        })
    return rows


def load_error_items():
    by_code, idx = load_master()
    rows = []
    for kcode, (sib_code, sib_name) in SIBLING_REF.items():
        row = by_code[kcode]
        list_price, cur_price = row[idx["기존단가"]], row[idx["최종인하단가"]]
        offer_price = by_code[sib_code][idx["최종인하단가"]]
        rows.append({
            "K코드": kcode, "품명": row[idx["품명"]].strip(), "마진분류": "-",
            "마진율": None, "정가": list_price, "대성_현재제시단가": cur_price,
            "대성_현재할인율": (list_price - cur_price) / list_price if list_price else None,
            "KT_제시매입단가": offer_price, "차액": cur_price - offer_price,
            "구분": "표기오류", "근거": (f"정가가 동일한 형제 품목 {sib_name}({sib_code})의 "
                                    f"실제 적용 단가를 기준으로 제시"),
        })
    return rows


def build_message_sheet(wb, margin_rows, err_rows):
    ws = wb.create_sheet("매입단가_제시(초안)")
    ws.append(["항목", "내용"])

    by_tier = {}
    for r in margin_rows:
        by_tier.setdefault(round(r["마진율"], 4), []).append(r)

    def tier_block(margin, label):
        items = sorted(by_tier.get(margin, []), key=lambda r: -r["정가"])
        if not items:
            return ""
        cur_rate = items[0]["대성_현재할인율"]
        example = items[0]
        return (
            f"{label} {len(items)}개 품목\n"
            f"현재 귀사 단가표상 할인율은 {cur_rate:.2%}입니다. 저희는 이 품목들을 귀사 "
            f"정가 그대로 고객에게 판매할 계획이며, 그 안에서 저희 통상 마진 {margin:.0%}를 "
            f"확보하려면 할인율 {margin:.0%} 기준 매입이 필요합니다. (예: "
            f"{example['품명']}({example['K코드']}) {example['대성_현재제시단가']:,.0f}원 "
            f"-> {example['KT_제시매입단가']:,.0f}원)\n"
            f"품목별 상세는 첨부 '요청단가_상세내역' 시트를 참고 부탁드립니다.\n\n"
        )

    err_lines = "\n".join(
        f"  - {r['품명']}({r['K코드']}): {r['KT_제시매입단가']:,.0f}원 (정가 {r['정가']:,}원)"
        for r in err_rows
    )

    msg = (
        "안녕하세요, KT commerce 이돈현입니다.\n\n"
        f"귀사 8GHz/11GHz 구성 단가표를 검토한 결과, 아래 {len(margin_rows) + len(err_rows)}개 "
        "품목은 매입단가를 다음과 같이 적용하여 진행하고자 합니다.\n\n"
        + tier_block(0.03, "1) 안테나/마운트류")
        + tier_block(0.05, "2) SFP")
        + tier_block(0.02, "3) 기타 핵심장비")
        + "4) RFUC-CPLR-8, RFUC-TWIST Kit-6\n"
        "두 품목 모두 할인 후 단가가 정가보다 높게(634,418원, 두 품목 동일값) 기재되어 "
        "있어 표기 오류로 판단됩니다. 정가가 동일한 형제 품목의 실제 적용 단가를 기준으로 "
        "아래와 같이 매입단가를 적용하겠습니다.\n"
        + err_lines + "\n\n"
        "위 품목을 반영한 전체 493개 품목 매입단가 목록을 첨부(전체493종_매입단가_제시 "
        "시트)합니다. 나머지 품목은 귀사가 제시하신 단가를 그대로 적용합니다.\n\n"
        "이견 있으시면 회신 부탁드리며, 별도 회신 없을 시 위 단가로 진행하겠습니다.\n\n"
        "감사합니다."
    )

    guide = [
        ("용도", "대성인포텍에 보낼 이메일/메신저 본문 초안. '요청단가_상세내역'과 "
                "'전체493종_매입단가_제시' 시트를 첨부하거나 표를 붙여넣으면 된다."),
        ("메시지 초안", msg),
        ("근거 설계(내부용, 대성에 보내지 않음)", "대성 자체 할인율이 정상/비정상인지는 "
                "따지지 않는다(안테나 2%는 대성 단가표 내 41개 동일 계열 전부와 일치하는 "
                "정상값). 대신 'KT는 정가 그대로 고객에게 판매하고, 그 안에서 통상 마진율을 "
                "남기려면 이 가격에 매입해야 한다'는 KT 자체 마진구조 논리로만 요청한다. "
                "안테나 36개는 2%->3%로 1%p만 올려달라는 요청이라 현실적이고, SFP는 "
                "1%->5%(4%p)로 이전 버전의 1%->14.69%(13.69%p) 요청보다 훨씬 작다."),
        ("주의", "이 초안은 자동 생성된 것이니 보내시기 전에 실제 상황(이미 통화/메일로 "
              "언급한 내용, 호칭, 마감 일정 등)에 맞게 다듬어서 사용하시기 바랍니다."),
    ]
    for k, v in guide:
        ws.append([k, v])
    finalize_sheet(ws, 1, 2, ws.max_row)
    ws.column_dimensions["B"].width = 100
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).alignment = WRAP_TOP


def build_list_sheet(wb, rows):
    ws = wb.create_sheet("요청단가_상세내역")
    headers = ["구분", "마진분류", "K코드", "품명", "정가", "대성 현재 제시단가",
               "대성 현재 할인율", "KT 제시 매입단가", "차액(현재-제시)", "근거"]
    ws.append(headers)
    for r in sorted(rows, key=lambda r: (r["구분"], r["마진분류"], -r["정가"])):
        rr = ws.max_row + 1
        ws.append([r["구분"], r["마진분류"], r["K코드"], r["품명"], r["정가"],
                   r["대성_현재제시단가"], r["대성_현재할인율"], r["KT_제시매입단가"],
                   r["차액"], r["근거"]])
        if r["구분"] == "표기오류":
            mark_fill(ws, rr, 1, ERROR_FILL)
            mark_fill(ws, rr, 7, ERROR_FILL)
            mark_fill(ws, rr, 8, CONFIRMED_FILL)
        else:
            mark_fill(ws, rr, 1, REVIEW_FILL)
            mark_fill(ws, rr, 8, REVIEW_FILL)
    for c in [5, 6, 8, 9]:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = FMT_AMOUNT
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=7).number_format = FMT_PERCENT
    finalize_sheet(ws, 1, len(headers), ws.max_row)
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["J"].width = 55
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=10).alignment = WRAP_TOP


def build_full_offer_sheet(wb, override_rows):
    by_code, idx = load_master()
    override = {r["K코드"]: r for r in override_rows}

    ws = wb.create_sheet("전체493종_매입단가_제시")
    headers = ["품목상태", "K코드", "품명", "정가(기존단가)", "대성 제시단가(최종인하단가)",
               "할인율", "KT 제시 매입단가", "조정구분", "비고"]
    ws.append(headers)

    n_accept = n_override = n_nolist = 0
    for kcode, row in by_code.items():
        status, name = row[idx["품목상태"]], row[idx["품명"]]
        list_price, cur_price = row[idx["기존단가"]], row[idx["최종인하단가"]]

        if kcode in override:
            o = override[kcode]
            adj = "마진 확보 위한 수정 제시" if o["구분"] == "마진확보" else "표기오류 수정 제시"
            offer, note = o["KT_제시매입단가"], o["근거"]
            rate = o["대성_현재할인율"]
            n_override += 1
        elif list_price is None or cur_price is None:
            offer, adj, note = cur_price, "정가 정보 없음 - 대성 제시가 그대로 수용", "신규 SW패키지 품목, 정가 기재 없어 자체 검증 불가"
            rate = None
            n_nolist += 1
        else:
            offer, adj, note = cur_price, "대성 제시가 수용", ""
            rate = (list_price - cur_price) / list_price if list_price else None
            n_accept += 1

        rr = ws.max_row + 1
        ws.append([status, kcode, name, list_price, cur_price, rate, offer, adj, note])
        if kcode in override:
            mark_fill(ws, rr, 7, CONFIRMED_FILL if override[kcode]["구분"] == "표기오류" else REVIEW_FILL)
            mark_fill(ws, rr, 8, REVIEW_FILL)
        elif status == "확인필요":
            mark_fill(ws, rr, 1, REVIEW_FILL)

    for c in [4, 5, 7]:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = FMT_AMOUNT
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=6).number_format = FMT_PERCENT
    finalize_sheet(ws, 1, len(headers), ws.max_row)
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["I"].width = 45
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=9).alignment = WRAP_TOP

    return n_accept, n_override, n_nolist


def main():
    margin_rows = load_margin_driven_items()
    err_rows = load_error_items()
    all_rows = margin_rows + err_rows

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_message_sheet(wb, margin_rows, err_rows)
    build_list_sheet(wb, all_rows)
    n_accept, n_override, n_nolist = build_full_offer_sheet(wb, all_rows)
    wb._sheets = [wb["매입단가_제시(초안)"], wb["요청단가_상세내역"], wb["전체493종_매입단가_제시"]]

    out_path = OUTPUT_DIR / "16_대성_매입단가_제시.xlsx"
    wb.save(out_path)

    total_down = sum(r["차액"] for r in all_rows)
    print("=== 16단계: 대성인포텍 매입단가 제시 목록 생성 완료 ===")
    print(f"마진확보 필요(정가상한 42건 중 실질 요청): {len(margin_rows)}건, "
          f"표기오류 추정(형제품목 대체): {len(err_rows)}건")
    print(f"현재 제시가 대비 인하 요청 총액: {total_down:,.0f}원")
    print(f"전체 493종 매입단가 제시: 대성 제시가 그대로 수용 {n_accept}건, 수정 제시 "
          f"{n_override}건, 정가 정보 없음(그대로 수용) {n_nolist}건")
    print(f"생성 파일: {out_path}")


if __name__ == "__main__":
    main()
