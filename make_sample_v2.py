"""
make_sample_v2.py — run_settlement.py 테스트용 샘플 데이터 생성
실제 KT/플랫폼 컬럼 구조 반영
  KT     : 구매문서번호, 구매품목, 납품일자, 공급가액, 공급자명, 세금코드명, 이동유형명, 인수증
  플랫폼 : 주문번호, 품목번호, 주문&품목, 입고일, 협력사명, 매출과세구분, 정산금액,
           매입금액, 매출총이익, 서비스카테고리, 중분류(IP만변경적용),
           일반/통신구분, 결제유형, 담당자, 담당자부서
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime

FONT = "Arial"
HDR_COLOR = "1F4E79"
HDR_GREEN = "375623"

def _bdr():
    s = Side(style="thin", color="C0C0C0")
    return Border(left=s, right=s, top=s, bottom=s)

def write_table(ws, headers, rows, header_color=HDR_COLOR, col_widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font      = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill      = PatternFill("solid", start_color=header_color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _bdr()
    ws.row_dimensions[1].height = 20
    for ri, row in enumerate(rows, 2):
        fill = "F2F2F2" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row, 1):
            cell = ws.cell(ri, ci, val)
            cell.font      = Font(name=FONT, size=10)
            cell.fill      = PatternFill("solid", start_color=fill)
            cell.border    = _bdr()
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(val, (int, float)) and not isinstance(val, bool) and abs(val) >= 1000:
                cell.number_format = "#,##0"
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

def d(day): return datetime.date(2025, 5, day)

# ══════════════════════════════════════════════════════
# 매핑 데이터 (샘플 전체에서 공통 사용)
# ══════════════════════════════════════════════════════

# 서비스카테고리 → (관리회계 카테고리, 담당자, 담당자부서)
SVC_MAP = {
    "출판/인쇄"  : ("출판/인쇄",   "정도원", "사무용품사업팀"),
    "IT서비스"   : ("IT서비스",    "김철수", "IT사업팀"),
    "통신장비"   : ("통신장비",    "이영희", "통신장비사업팀"),
    "소프트웨어" : ("소프트웨어",  "박민준", "소프트웨어사업팀"),
    "유지보수"   : ("유지보수",    "최수진", "유지보수사업팀"),
}
# 협력사 → 서비스카테고리
CORP_SVC = {
    "우리문구(주)"    : "출판/인쇄",
    "한국IT솔루션(주)": "IT서비스",
    "케이네트웍스(주)": "통신장비",
    "비즈솔루션(주)"  : "소프트웨어",
    "에코서비스(주)"  : "유지보수",
    "삼성SDS"         : "IT서비스",
}
# 담당자 → 일반/통신구분
TYPE_MAP = {"정도원": "일반", "김철수": "일반", "박민준": "일반",
            "이영희": "통신", "최수진": "통신"}


# ══════════════════════════════════════════════════════
# 1. mapping_master.xlsx
#    실제 컬럼: 카테고리 UID | 카테고리명 | 서비스 카테고리 |
#               관리회계 카테고리(안) | 담당부서(안) | 상품담당자(26.05)
# ══════════════════════════════════════════════════════
wb_mm = openpyxl.Workbook()

ws1 = wb_mm.active; ws1.title = "서비스카테고리"
write_table(ws1,
    ["카테고리 UID", "카테고리명", "서비스 카테고리",
     "관리회계 카테고리(안)", "담당부서(안)", "상품담당자(26.05)"],
    [
        (1001, "인쇄물",    "출판/인쇄",   "출판/인쇄",   "사무용품사업팀",   "정도원"),
        (1002, "IT서비스",  "IT서비스",    "IT서비스",    "IT사업팀",         "김철수"),
        (1003, "통신장비",  "통신장비",    "통신장비",    "통신장비사업팀",   "이영희"),
        (1004, "소프트웨어","소프트웨어",  "소프트웨어",  "소프트웨어사업팀", "박민준"),
        (1005, "유지보수",  "유지보수",    "유지보수",    "유지보수사업팀",   "최수진"),
    ],
    col_widths=[12, 12, 14, 18, 18, 16]
)

ws2 = wb_mm.create_sheet("일반통신구분")
write_table(ws2,
    ["담당자", "일반/통신구분"],
    [
        ("정도원", "일반"), ("김철수", "일반"), ("박민준", "일반"),
        ("이영희", "통신"), ("최수진", "통신"),
    ],
    col_widths=[14, 16]
)

ws3 = wb_mm.create_sheet("기업규모")
write_table(ws3,
    ["협력사명", "기업규모"],
    [
        ("삼성SDS",           "대기업"),
        ("한국IT솔루션(주)",  "중견기업"),
        ("우리문구(주)",      "중소기업"),
        ("케이네트웍스(주)",  "중소기업"),
        ("비즈솔루션(주)",    "중소기업"),
        ("에코서비스(주)",    "중소기업"),
    ],
    col_widths=[20, 14]
)
wb_mm.save("mapping_master.xlsx")
print("[OK] mapping_master.xlsx 생성 (실제 컬럼명 적용)")

# ══════════════════════════════════════════════════════
# 2. KT raw 데이터 (실제 KT 컬럼 구조)
# ══════════════════════════════════════════════════════
TAX_CODE = {
    "과세":   "과세표준세율(10%)",
    "면세":   "면세(별도세액없음)",
    "비과세": "비과세",
}

KT_HEADERS = [
    "입고전표", "구매문서번호", "구매품목",
    "납품일자", "공급자명",
    "세금코드명", "이동유형명", "인수증",
    "공급가액", "세액", "합계금액", "사업장",
]

def kt_row(doc_no, item_no, day, corp, tax_key, move, req_no, supply_amt, biz="본사"):
    tax_code = TAX_CODE[tax_key]
    vat = supply_amt * 0.1 if tax_key == "과세" else 0
    total = supply_amt + int(vat)
    return (
        f"EL{doc_no}",   # 입고전표
        doc_no,           # 구매문서번호
        item_no,          # 구매품목 (복합 키 후반부)
        d(day),           # 납품일자
        corp,             # 공급자명
        tax_code,         # 세금코드명 → _preprocess_kt()가 정규화
        move,             # 이동유형명: '입고' or '반품'
        req_no,           # 인수증
        supply_amt,       # 공급가액
        int(vat),         # 세액
        total,            # 합계금액
        biz,              # 사업장
    )

kt_rows = []

# ─ A. 정상 일치 15건 ─
normal_data = [
    # (주문번호, 구매품목, 요청번호, 입고일, 협력사, 과세, 금액)
    # 주문번호 ORD-2025-001 에 2개 라인 (10, 20)
    ("ORD-2025-001", "10", "REQ-001",  1, "우리문구(주)",     "과세",     300_000),
    ("ORD-2025-001", "20", "REQ-001",  1, "우리문구(주)",     "과세",     200_000),
    # 이하 주문번호당 1개 라인 (10)
    ("ORD-2025-002", "10", "REQ-002",  1, "우리문구(주)",     "과세",     150_000),
    ("ORD-2025-003", "10", "REQ-003",  2, "우리문구(주)",     "과세",     500_000),
    ("ORD-2025-004", "10", "REQ-004",  2, "한국IT솔루션(주)", "과세",   1_200_000),
    ("ORD-2025-005", "10", "REQ-005",  3, "한국IT솔루션(주)", "면세",     850_000),
    ("ORD-2025-006", "10", "REQ-006",  5, "케이네트웍스(주)", "과세",   7_000_000),
    ("ORD-2025-007", "10", "REQ-007",  6, "케이네트웍스(주)", "과세",   5_200_000),
    ("ORD-2025-008", "10", "REQ-008",  7, "비즈솔루션(주)",   "과세",   8_900_000),
    ("ORD-2025-009", "10", "REQ-009",  8, "비즈솔루션(주)",   "과세",   6_300_000),
    ("ORD-2025-010", "10", "REQ-010",  9, "에코서비스(주)",   "면세",     350_000),
    ("ORD-2025-011", "10", "REQ-011", 10, "에코서비스(주)",   "면세",     840_000),
    ("ORD-2025-012", "10", "REQ-012", 12, "우리문구(주)",     "과세",      80_000),
    ("ORD-2025-013", "10", "REQ-013", 13, "한국IT솔루션(주)", "과세",   2_500_000),
    ("ORD-2025-014", "10", "REQ-014", 14, "케이네트웍스(주)", "과세",   3_400_000),
]
for doc, item, req, day, corp, tax, amt in normal_data:
    kt_rows.append(kt_row(doc, item, day, corp, tax, "입고", req, amt))

# ─ B. 플랫폼 누락 3건 (KT에만 있음) ─
missing_data = [
    ("ORD-2025-015", "10", "REQ-015",  4, "우리문구(주)",     "과세",     225_000),
    ("ORD-2025-016", "10", "REQ-016",  8, "한국IT솔루션(주)", "과세",     380_000),
    ("ORD-2025-017", "10", "REQ-017", 11, "에코서비스(주)",   "면세",     150_000),
]
for doc, item, req, day, corp, tax, amt in missing_data:
    kt_rows.append(kt_row(doc, item, day, corp, tax, "입고", req, amt))

# ─ C. 반품 2건 (이동유형명='반품') ─
return_data = [
    ("RTN-2025-001", "10", "REQ-018",  6, "케이네트웍스(주)", "과세", -3_500_000),
    ("RTN-2025-002", "10", "REQ-019", 13, "에코서비스(주)",   "면세",   -420_000),
]
for doc, item, req, day, corp, tax, amt in return_data:
    kt_rows.append(kt_row(doc, item, day, corp, tax, "반품", req, amt))

# ─ D. 금액 불일치 5건 (KT = 플랫폼 + 500) ─
amtdiff_data = [
    ("ORD-2025-018", "10", "REQ-020",  3, "비즈솔루션(주)",   "과세",  1_750_500),
    ("ORD-2025-019", "10", "REQ-021",  5, "한국IT솔루션(주)", "과세",  2_300_500),
    ("ORD-2025-020", "10", "REQ-022",  9, "케이네트웍스(주)", "과세",  3_100_500),
    ("ORD-2025-021", "10", "REQ-023", 12, "우리문구(주)",     "과세",    980_500),
    ("ORD-2025-022", "10", "REQ-024", 14, "비즈솔루션(주)",   "과세",  4_500_500),
]
for doc, item, req, day, corp, tax, amt in amtdiff_data:
    kt_rows.append(kt_row(doc, item, day, corp, tax, "입고", req, amt))

# ─ E. 비과세 대기업 1건 (어음 테스트) ─
kt_rows.append(kt_row("ORD-2025-023", "10", 10, "삼성SDS", "비과세", "입고", "REQ-025", 210_000_000))

# 합계 행 (구매문서번호 = '합계' → _preprocess_kt()가 자동 제거)
total_supply = sum(r[8] for r in kt_rows)
kt_rows.append((
    "", "합계", "", "", "", "", "", "",
    total_supply, "", "", "",
))

wb_kt = openpyxl.Workbook()
ws_kt = wb_kt.active; ws_kt.title = "KT_입고내역"
write_table(ws_kt, KT_HEADERS, kt_rows,
            col_widths=[14, 18, 10, 13, 20, 22, 12, 14, 16, 12, 16, 8])
wb_kt.save("kt_raw.xlsx")
print(f"[OK] kt_raw.xlsx 생성 ({len(kt_rows)-1}건 + 합계행 1건)")

# ══════════════════════════════════════════════════════
# 3. platform.xlsx  (실제 플랫폼 컬럼 구조)
#    ★ 주문&품목 = 주문번호+품목번호 (구분자 없이 연결)
#    ★ 중분류(IP만변경적용) = 공란 (run_settlement.py가 XLOOKUP으로 채움)
#    ★ 매입금액, 매출총이익 추가
# ══════════════════════════════════════════════════════

PL_HEADERS = [
    "주문번호", "품목번호", "주문&품목",
    "입고일", "협력사명", "매출과세구분",
    "정산금액", "매입금액", "매출총이익",
    "서비스카테고리", "중분류(IP만변경적용)",
    "일반/통신구분", "결제유형", "담당자", "담당자부서",
]

def buy_amt(amt, tax_key):
    """매입금액: 과세 92%, 면세 95%, 비과세 90%"""
    ratio = {"과세": 0.92, "면세": 0.95, "비과세": 0.90}.get(tax_key, 0.92)
    return int(amt * ratio)

def pl_row(doc, item, day, corp, tax, amt):
    svc       = CORP_SVC.get(corp, "")
    acc, mgr, dept = SVC_MAP.get(svc, ("", "", ""))
    tp        = TYPE_MAP.get(mgr, "")
    combo_key = f"{doc}{item}"         # 주문&품목 (문자열 연결)
    b_amt     = buy_amt(amt, tax)
    gp_amt    = amt - b_amt
    return (
        doc, item, combo_key,
        d(day), corp, tax,
        amt, b_amt, gp_amt,
        svc, "",           # 중분류(IP만변경적용) = 공란 (XLOOKUP으로 채워짐)
        tp, "현금", mgr, dept,
    )

pl_rows = []

# 정상 15건
for doc, item, req, day, corp, tax, amt in normal_data:
    pl_rows.append(pl_row(doc, item, day, corp, tax, amt))

# 금액 불일치 5건 (플랫폼 금액 = KT - 500)
for doc, item, req, day, corp, tax, amt in amtdiff_data:
    pl_rows.append(pl_row(doc, item, day, corp, tax, amt - 500))

# 반품 원주문 2건 (플랫폼 별도 주문번호로 기록됨)
pl_rows.append(pl_row("ORD-ORIG-018", "10",  5, "케이네트웍스(주)", "과세", -3_500_000))
pl_rows.append(pl_row("ORD-ORIG-019", "10", 11, "에코서비스(주)",   "면세",   -420_000))

# 대기업 비과세 1건
pl_rows.append(pl_row("ORD-2025-023", "10", 10, "삼성SDS", "비과세", 210_000_000))

wb_pl = openpyxl.Workbook()
ws_pl = wb_pl.active; ws_pl.title = "플랫폼_입고내역"
write_table(ws_pl, PL_HEADERS, pl_rows,
            col_widths=[18, 10, 20, 13, 20, 14, 16, 16, 16, 14, 20, 14, 10, 12, 18])
wb_pl.save("platform.xlsx")
print(f"[OK] platform.xlsx 생성 ({len(pl_rows)}건, 요청번호 없음 → 경로 B)")

# ══════════════════════════════════════════════════════
# 4. return_mapping.xlsx
#    KT반품복합키  = 구매문서번호 + 구매품목 (구분자 없이)
#    플랫폼원복합키 = 플랫폼주문번호 + 품목번호 (구분자 없이, 사용자 직접 입력)
# ══════════════════════════════════════════════════════
RM_HEADERS = [
    "KT반품복합키", "KT반품주문번호", "KT구매품목", "KT요청번호",
    "플랫폼원복합키", "협력사명", "정산금액", "처리상태",
]
rm_rows = [
    # KT반품복합키 = RTN-2025-001 + 10 = RTN-2025-00110
    # 플랫폼원복합키 = ORD-ORIG-018 + 10 = ORD-ORIG-01810
    ("RTN-2025-00110", "RTN-2025-001", "10", "REQ-018",
     "ORD-ORIG-01810", "케이네트웍스(주)", -3_500_000, "완료"),
    ("RTN-2025-00210", "RTN-2025-002", "10", "REQ-019",
     "ORD-ORIG-01910", "에코서비스(주)",     -420_000, "완료"),
]
wb_rm = openpyxl.Workbook()
ws_rm = wb_rm.active; ws_rm.title = "반품매핑"
write_table(ws_rm, RM_HEADERS, rm_rows,
            col_widths=[18, 16, 10, 14, 18, 20, 16, 10])
wb_rm.save("return_mapping.xlsx")
print("[OK] return_mapping.xlsx 생성 (플랫폼원복합키 입력 완료)")

# ══════════════════════════════════════════════════════
# 요약
# ══════════════════════════════════════════════════════
통신_corps = [c for c, s in CORP_SVC.items() if SVC_MAP.get(s, ("","",""))[2] == "통신"]
일반_corps = [c for c, s in CORP_SVC.items() if SVC_MAP.get(s, ("","",""))[2] == "일반"]
print(f"""
[샘플 데이터 요약]
  kt_raw.xlsx       : {len(kt_rows)-1}건 (+ 합계 행 1건, _preprocess_kt() 자동 제거)
    - 정상 일치     : 15건  (ORD-2025-001 ~ 014, 001에 2라인)
    - 플랫폼 누락   :  3건  (ORD-2025-015 ~ 017)
    - 반품           :  2건  (RTN-2025-001 ~ 002, 이동유형명='반품')
    - 금액 불일치   :  5건  (ORD-2025-018 ~ 022, KT+500)
    - 대기업/비과세 :  1건  (ORD-2025-023, 삼성SDS 2.1억)

  platform.xlsx     : {len(pl_rows)}건 (요청번호 없음 → 경로 B)
    - 정상 15건 + 금액차이 5건 + 반품원주문 2건 + 대기업 1건
    - 누락 3건 미포함
    ★ 신규 컬럼: 주문&품목, 매입금액, 매출총이익, 중분류(IP만변경적용)=공란,
                  일반/통신구분, 결제유형, 담당자, 담당자부서

  매칭 키: 구매문서번호+구매품목 (KT) ↔ 주문&품목 (플랫폼, 사전계산)

  mapping_master.xlsx:
    시트1 실제 컬럼: 카테고리 UID | 카테고리명 | 서비스 카테고리 |
                      관리회계 카테고리(안) | 담당부서(안) | 상품담당자(26.05)
    서비스카테고리: {len(SVC_MAP)}건 / 일반통신구분: 5건 / 기업규모: 6건

  통신 협력사: {통신_corps}
  일반 협력사: {일반_corps}

  return_mapping.xlsx: 반품 2건 (플랫폼원복합키 입력 완료)
""")
